"""Visible per-card material tables with recoverable Word-contained quantities."""
import json
import re
from xml.etree import ElementTree as ET

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Cm, Pt

from lab_quantities import quantities_hash, table_matrix

NAMESPACE = 'urn:reportskills:materials:v1'


def add_material_table(cell, rows):
    matrix = table_matrix(rows)
    cell.add_paragraph('用药与样品用量')
    table = cell.add_table(rows=1, cols=5)
    table.autofit = False
    widths = [3.1, 2.0, 2.6, 3.1, 6.2]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    table.style = 'Table Grid'
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for idx, values in enumerate(matrix):
        cells = table.rows[0].cells if idx == 0 else table.add_row().cells
        for target, text, width in zip(cells, values, widths):
            target.width = Cm(width)
            target.text = text
            for p in target.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.bold = idx == 0
    return quantities_hash([[c.text for c in r.cells] for r in table.rows])


def attach_material_payload(doc, records):
    if not records:
        return
    root = ET.Element('{' + NAMESPACE + '}materials')
    for record_id, value in records.items():
        node = ET.SubElement(root, '{' + NAMESPACE + '}record', {'id': record_id})
        node.text = json.dumps(value, ensure_ascii=False, allow_nan=False)
    blob = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    part = Part(PackURI('/customXml/reportskills-materials.xml'), 'application/xml', blob, doc.part.package)
    doc.part.relate_to(part, RT.CUSTOM_XML)


def read_material_payload(doc):
    payload = {}
    for relation in doc.part.rels.values():
        if relation.reltype != RT.CUSTOM_XML or relation.is_external:
            continue
        try:
            root = ET.fromstring(relation.target_part.blob)
        except ET.ParseError:
            continue
        if root.tag != '{' + NAMESPACE + '}materials':
            continue
        for node in root:
            if node.get('id') in payload:
                raise ValueError('Word 中重复的用药数据记录标识')
            payload[node.get('id')] = json.loads(node.text)
    matched = set()
    for outer in doc.tables:
        for row in outer.rows:
            for cell in row.cells:
                match = re.search(r'【记录标识】([^\n]+)', cell.text)
                if not match or match[1].strip() not in payload:
                    continue
                rid = match[1].strip()
                if rid in matched:
                    continue
                expected = payload[rid]
                if len(cell.tables) != 1:
                    raise ValueError('用药表被删除或结构已修改：' + rid + '；请定点核对后更新结构化用量')
                matrix = [[c.text for c in r.cells] for r in cell.tables[0].rows]
                if quantities_hash(matrix) != expected['table_hash']:
                    raise ValueError('用药表存在手动修改：' + rid + '；原始表格与结构化数据需先协调，不能静默覆盖')
                matched.add(rid)
    if matched != set(payload):
        raise ValueError('Word 用药表与记录标识不匹配')
    return {rid: value['materials'] for rid, value in payload.items()}
