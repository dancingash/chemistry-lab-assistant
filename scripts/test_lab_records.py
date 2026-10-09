"""Behavioral regressions; all Word writes are isolated from production reports."""
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from extract_daily_docx import extract_docx
from lab_records import Store, Conflict, tracker_name, tracker_data, write_daily
from project_context import update
import yaml
from build_monthly_summary import detect_project_ids, fallback_content, main as build_monthly_main


def record(rid='r1', day='2026-09-17', sample='DEMO-20260917-1', **kw):
    data = dict(record_id=rid, date=day, created_date=day, project='DEMO 示例项目',
                sample_id=sample, step='S01', experiment_type='合成', state='合成中',
                today='实验用合成记录。', details='原料用量：待补。', follow_up='完成后处理。')
    data.update(kw)
    return data


class RecordsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.store = Store(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def tracker(self, month='2026-09'):
        return self.root/'report'/month/tracker_name(month)

    def test_save_retry_and_numeric_sort(self):
        self.store.save([record(), record('r2', sample='DEMO-20260917-10'), record('r3', sample='DEMO-20260917-2')])
        self.store.save([record()])
        self.assertEqual(len(self.store.query(start='2026-09-17')), 3)
        rows = tracker_data(self.tracker())['rows']
        self.assertEqual([r[0] for r in rows], ['DEMO-20260917-10','DEMO-20260917-2','DEMO-20260917-1'])
        self.store.save([record('r4','2026-09-18', created_date='2026-09-17', experiment_type='后处理', state='待测试')])
        self.assertEqual(tracker_data(self.tracker())['rows'][-1][2], '2026-09-17')
        self.assertEqual(len(self.store.query(sample='DEMO-20260917-1')), 2)

    def test_rollover_only_unfinished_and_archive(self):
        self.store.save([record(), record('r2', sample='DEMO-20260917-2', state='已完成')])
        original = self.tracker().read_bytes()
        self.store.write_tracker('2026-10')
        self.assertEqual(self.tracker().read_bytes(), original)
        self.assertEqual(len(tracker_data(self.tracker('2026-10'))['rows']), 1)
        self.store.save([record('r3','2026-10-02','DEMO-20261002-1')])
        self.assertEqual(len(tracker_data(self.tracker('2026-10'))['rows']), 2)
        with self.assertRaises(Conflict):
            self.store.write_tracker('2026-09')

    def test_branch_and_older_step_test(self):
        self.store.save([record()])
        self.store.save([record('r2', '2026-09-18', created_date='2026-09-17', step='S02')])
        self.store.save([record('r3','2026-09-19',created_date='2026-09-17',experiment_type='测试',state='已完成')])
        row = tracker_data(self.tracker())['rows'][0]
        self.assertEqual(row[3], 'S02 / 合成中')
        branch = record('r4','2026-09-19','DEMO-20260917-1-B01',source_sample='DEMO-20260917-1',source_step='S02')
        self.store.save([branch])
        found = self.store.query(sample=branch['sample_id'])
        self.assertEqual(found[0]['source_step'], 'S02')
        self.assertEqual(len(self.store.query(sample='DEMO-20260917-1',step='S01',kind='测试')),1)

    def test_selective_query_and_manual_daily_edit(self):
        self.store.save([record()])
        self.store.save([record('r2','2026-09-18','DEMO-20260918-1')])
        store = Store(self.root); store.refresh(); store.log.clear()
        selected = store.query(sample='DEMO-20260917-1')
        self.assertEqual(len(selected),1)
        self.assertEqual(len(store.log),1)
        self.assertIn('2026-09-17.docx', store.log[0]['path'])
        path = self.root/'report/2026-09/日报_2026-09-17.docx'
        doc = Document(path)
        cell = doc.tables[0].cell(0,0)
        cell.text = cell.text.replace('原料用量：待补。','原料用量：实测 1.23 g。')
        doc.save(path)
        store = Store(self.root); store.refresh()
        self.assertEqual(len(store.log),1)
        self.assertIn('1.23',store.query(sample='DEMO-20260917-1')[0]['details'])

    def test_tracker_manual_conflict_preserved(self):
        self.store.save([record()])
        path = self.tracker()
        doc=Document(path); doc.tables[0].cell(1,3).text='S01 / 已完成'; doc.save(path)
        before=path.read_bytes()
        with self.assertRaises(Conflict):
            Store(self.root).refresh()
        self.assertEqual(before,path.read_bytes())

    def test_rebuild_and_interruption_retry(self):
        self.store.save([record()])
        (self.root/'report/index.json').unlink()
        recovered=Store(self.root); recovered.refresh()
        self.assertEqual(len(recovered.query(sample='DEMO-20260917-1')),1)
        with patch.object(recovered,'write_tracker',side_effect=OSError('simulated interruption')):
            with self.assertRaises(OSError):
                recovered.save([record('r2','2026-09-18','DEMO-20260918-1')])
        resumed=Store(self.root)
        resumed.save([record('r2','2026-09-18','DEMO-20260918-1')])
        self.assertEqual(len(resumed.query(start='2026-09-17')),2)
        self.assertEqual(len(tracker_data(self.tracker())['rows']),2)

    def test_legacy_paragraphs_spaces_and_snapshots(self):
        folder=self.root/'report/2026-08'; folder.mkdir(parents=True)
        a=folder/'old-a.docx'; b=folder/'old-b.docx'
        doc=Document(); doc.add_paragraph('2026 年 08 月 01 日 · 周六')
        for sample in ('DEMO-20260801-1','原编号未提供'):
            for line in [f'【项目】DEMO 示例',f'【样品编号】{sample}','【实验类型】测试','【配方与细节】实测 1 g','【后续跟进】待补']:
                doc.add_paragraph(line)
        doc.save(a); doc.save(b)
        parsed=extract_docx(a)['records']; self.assertEqual(len(parsed),2)
        self.assertEqual(parsed[0]['date'],'2026-08-01')
        self.assertEqual(parsed[1]['sample_ids'],[])
        self.assertEqual(parsed[0]['record_id'],extract_docx(b)['records'][0]['record_id'])
        self.store.refresh([a.relative_to(self.root)])
        with self.assertRaises(Conflict):
            self.store.refresh([b.relative_to(self.root)])

    def test_duplicate_ids_and_manual_unknown_content(self):
        self.store.save([record()])
        path=self.root/'report/2026-09/日报_2026-09-18.docx'
        write_daily(path,[record(day='2026-09-18')])
        with self.assertRaises(Conflict): Store(self.root).refresh()
        path.unlink()
        path=self.root/'report/2026-09/日报_2026-09-17.docx'
        doc=Document(path); doc.add_paragraph('用户手写备注'); doc.save(path)
        with self.assertRaises(Conflict): Store(self.root).save([record()])
        self.assertIn('用户手写备注','\n'.join(p.text for p in Document(path).paragraphs))

    def test_cross_month_query(self):
        self.store.save([record(day='2026-09-30',sample='DEMO-20260930-1')])
        self.store.save([record('r2','2026-10-01','DEMO-20261001-1')])
        self.assertEqual(len(self.store.query(start='2026-09-30',end='2026-10-01')),2)

    def test_monthly_identity_and_untruncated_details(self):
        r = record(step='S02', details='实测用量 1.23 g；' * 30)
        self.assertEqual(detect_project_ids([r]), ['DEMO'])
        self.assertEqual(detect_project_ids([dict(r,project='DEMO_SUB 项目',sample_id='DEMO_SUB-20260917-1')]),['DEMO_SUB'])
        content=fallback_content([r])
        self.assertIn('S02',content)
        self.assertIn(r['details'],content)

    def test_monthly_legacy_mapping_starts_on_effective_date(self):
        references = self.root/'references'
        references.mkdir()
        for project_id, extra in (
            ('DEMO_MAIN', {'effective_date': '2026-09-18', 'legacy_project_ids': ['DEMO_OLD']}),
            ('DEMO_OTHER', {}), ('DEMO_OLD', {}),
        ):
            data = dict(project_id=project_id, project_name=project_id, status='active',
                        targets={'BET': dict(type='qualitative', requirement='待确定',
                                             target_status='pending', required=True)})
            data.update(extra)
            (references/f'{project_id}.yaml').write_text(yaml.safe_dump(data, allow_unicode=True), encoding='utf-8')
        records = [
            record('old', '2026-09-17', 'DEMO_OLD-20260917-1', project='DEMO_OLD 前体'),
            record('mapped', '2026-09-18', 'DEMO_OLD-20260918-1', project='DEMO_OLD 前体'),
            record('separate', '2026-09-18', 'DEMO_OTHER-20260918-1', project='DEMO_OTHER 低氯树脂'),
        ]
        source = self.root/'records.json'
        output = self.root/'summary.json'
        source.write_text(json.dumps({'source': 'test', 'records': records}, ensure_ascii=False), encoding='utf-8')
        with patch.object(sys, 'argv', ['build_monthly_summary', str(source), '--references-dir',
                                        str(references), '--out', str(output), '--month', '2026-09']):
            build_monthly_main()
        rows = json.loads(output.read_text(encoding='utf-8'))['rows']
        by_project = {row['project'].split()[0]: row for row in rows}
        self.assertEqual(set(by_project), {'DEMO_MAIN', 'DEMO_OTHER', 'DEMO_OLD'})
        self.assertIn('DEMO_OLD-20260918-1', by_project['DEMO_MAIN']['content'])
        self.assertNotIn('DEMO_OLD-20260917-1', by_project['DEMO_MAIN']['content'])
        self.assertIn('DEMO_OLD-20260917-1', by_project['DEMO_OLD']['content'])
        self.assertIn('DEMO_OTHER-20260918-1', by_project['DEMO_OTHER']['content'])
        self.assertEqual(by_project['DEMO_MAIN']['target'], 'BET：数值目标待确定')

    def test_card_styles_and_navigation(self):
        self.store.save([record(), record('r2',sample='DEMO-20260917-2',experiment_type='后处理'),
                         record('r3',sample='DEMO-20260917-3',experiment_type='测试',today='长记录。'*1200)])
        doc=Document(self.root/'report/2026-09/日报_2026-09-17.docx')
        self.assertEqual(len(doc.tables),3)
        for table,fill in zip(doc.tables,('EFF6FF','F0FDF4','FFF7ED')):
            self.assertIn(fill,table._tbl.xml)
            self.assertNotIn('cantSplit',table._tbl.xml)
        tracker=Document(self.tracker())
        links=[rel for rel in tracker.part.rels.values() if rel.reltype.endswith('/hyperlink')]
        self.assertTrue(links)
        for rel in links:
            self.assertTrue((self.tracker().parent/rel.target_ref).exists())

    def test_migrate_uncertain_legacy_day_and_append(self):
        folder=self.root/'report/2026-09'; folder.mkdir(parents=True)
        old=folder/'日报_2026年09月_updated_0917.docx'
        doc=Document()
        for line in ['2026 年 09 月 17 日 · 周四','【项目】DEMO 示例项目',
                     '【样品编号】昨日树脂（原编号未提供）','【实验类型】后处理与干燥',
                     '【今日实验】已洗涤。','【配方与细节】用量待补。','【后续跟进】安排测试。']:
            doc.add_paragraph(line)
        doc.save(old)
        self.store.refresh([old.relative_to(self.root)])
        card=self.store.migrate_day('2026-09-17',old.relative_to(self.root))
        self.assertEqual(len(self.store.query(start='2026-09-17',end='2026-09-17')),1)
        self.assertIn('【来源记录标识】',Document(card).tables[0].cell(0,0).text)
        self.assertTrue(self.tracker().exists())
        self.store.save([record('new-record','2026-09-17','DEMO-20260917-1')])
        self.assertEqual(len(self.store.query(start='2026-09-17',end='2026-09-17')),2)
        self.assertEqual(len(Document(card).tables),2)
        self.assertEqual(self.store.migrate_day('2026-09-17',old.relative_to(self.root)),card)

    def test_fresh_note_without_sample_id_is_still_a_card(self):
        unresolved=record(sample='昨日树脂（原编号未提供）',created_date='待确认',step='待确认',state='待确认')
        path=self.store.save([unresolved])
        self.assertTrue(path.exists())
        self.assertEqual(tracker_data(self.tracker())['rows'],[])
        self.assertEqual(len(self.store.query(start='2026-09-17',end='2026-09-17')),1)
        corrected=dict(unresolved,sample_id='DEMO-20260916-1',created_date='2026-09-16',step='S01',state='待测试')
        self.store.save([corrected])
        self.assertEqual(len(Document(path).tables),1)
        self.assertEqual(tracker_data(self.tracker())['rows'][0][0],'DEMO-20260916-1')

    def test_context_history_preserves_custom_fields(self):
        candidate=self.root/'candidate.yaml'
        background=dict(project_id='DEMO',project_name='示例',status='active',project_route='待补',
                        targets={'BET':dict(type='min',value=10,unit='m²/g')},chart_groups={},custom='preserve',reference_sources=[])
        candidate.write_text(yaml.safe_dump(background,allow_unicode=True),encoding='utf-8')
        target=update(self.root,candidate,'2026-09-17')
        background['targets']['BET']['value']=20
        candidate.write_text(yaml.safe_dump(background,allow_unicode=True),encoding='utf-8')
        update(self.root,candidate,'2026-09-18')
        old=yaml.safe_load((target.parent/'history/DEMO/v0001.yaml').read_text(encoding='utf-8'))
        self.assertEqual(old['targets']['BET']['value'],10)
        self.assertEqual(yaml.safe_load(target.read_text(encoding='utf-8'))['context_version'],2)
        del background['custom']
        candidate.write_text(yaml.safe_dump(background),encoding='utf-8')
        with self.assertRaises(ValueError): update(self.root,candidate,'2026-09-19')


if __name__=='__main__':
    unittest.main()
