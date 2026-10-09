"""Version reusable experiment plans; explicit best selection, Word exports and daily links."""
from __future__ import annotations
from lab_config import data_root

import argparse
from contextlib import contextmanager
from copy import deepcopy
from datetime import date, datetime
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sys
import tempfile

SCHEMA = 1
MAP_SECTIONS = ('materials', 'equipment', 'operations', 'conditions', 'controls', 'tests', 'acceptance')
SECTIONS = ('purpose', 'scope', *MAP_SECTIONS, 'pending')
LABELS = {'purpose': '试验目的', 'materials': '原料配方', 'operations': '操作步骤',
          'conditions': '操作条件', 'tests': '测试项目', 'pending': '待补信息',
          'operation_order': '操作顺序'}
LABELS.update(scope='适用范围', equipment='设备', controls='关键控制点', acceptance='检验与判定要求')
SOURCE_FIELDS = {'description', 'location', 'record_id', 'sample_id', 'context_version',
                 'source_type', 'title', 'locator', 'original_method', 'adaptation_reason'}
PLAN_FIELDS = ('plan_stage', 'plan_id', 'plan_version')
IDENTIFIER = re.compile(r'[A-Za-z][A-Za-z0-9_-]{0,79}')
RESERVED = re.compile(r'(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])', re.I)


class PlanConflict(ValueError):
    pass


def now():
    return datetime.now().astimezone().isoformat(timespec='seconds')


def digest(data):
    return hashlib.sha256(data).hexdigest()


def json_bytes(value):
    return json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False).encode('utf-8')


def atomic_write(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=path.parent, suffix='.tmp')
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def read_json(path):
    try:
        return json.loads(path.read_text(encoding='utf-8-sig'))
    except (ValueError, OSError) as exc:
        raise PlanConflict(f'无法读取 {path}：{exc}') from exc


def identifier(value):
    if not isinstance(value, str) or not IDENTIFIER.fullmatch(value) or RESERVED.fullmatch(value):
        raise PlanConflict('项目及工序编号须为安全的字母数字编号，可含下划线或短横线')
    return value


def nonempty(value, label):
    if not isinstance(value, str) or not value.strip():
        raise PlanConflict(f'{label}须为非空文本')
    return value


def version_number(value):
    if isinstance(value, bool) or not re.fullmatch(r'[1-9]\d*', str(value)):
        raise PlanConflict('须指定正整数版本号，日报引用不能使用 best')
    return int(value)


def validate_content(content):
    if not isinstance(content, dict) or set(content) - set(SECTIONS):
        raise PlanConflict('未知方案内容字段')
    result = dict(purpose='待补', scope='待补', pending=[], **{key: {} for key in MAP_SECTIONS})
    result.update(deepcopy(content))
    nonempty(result['purpose'], '试验目的')
    nonempty(result['scope'], '适用范围')
    for section in MAP_SECTIONS:
        if not isinstance(result[section], dict):
            raise PlanConflict(f'{section} 须为名称到文本的映射')
        for key, value in result[section].items():
            nonempty(key, section + '名称')
            nonempty(value, section + '内容（包含单位或待补标记）')
    if not isinstance(result['pending'], list) or any(not isinstance(x, str) or not x.strip() for x in result['pending']):
        raise PlanConflict('pending 须为待补信息文本列表')
    return result


def merge_content(base, patch, operation_order=None):
    if not isinstance(patch, dict) or set(patch) - set(SECTIONS):
        raise PlanConflict('未知方案修改字段')
    result = deepcopy(base)
    for section, values in patch.items():
        if section in MAP_SECTIONS:
            if not isinstance(values, dict):
                raise PlanConflict('配方、步骤、条件和测试项目须逐项修改')
            for key, value in values.items():
                if value is None:
                    if key not in result.setdefault(section, {}):
                        raise PlanConflict(f'不能删除不存在的项目：{section}/{key}')
                    del result[section][key]
                else:
                    result.setdefault(section, {})[key] = value
        else:
            result[section] = deepcopy(values)
    if operation_order is not None:
        if (not isinstance(operation_order, list) or any(not isinstance(k, str) for k in operation_order)
                or len(operation_order) != len(set(operation_order))
                or set(operation_order) != set(result['operations'])):
            raise PlanConflict('operation_order 须恰好包含修改后每个操作步骤名称一次')
        result['operations'] = {key: result['operations'][key] for key in operation_order}
    return validate_content(result)


def changes_between(base, content):
    changes = []
    for section in SECTIONS:
        before, after = base.get(section), content[section]
        if section in MAP_SECTIONS:
            for key in dict.fromkeys([*(before or {}), *after]):
                old, new = (before or {}).get(key), after.get(key)
                if old != new:
                    changes.append({'field': [section, key], 'before': old, 'after': new})
        elif before != after:
            changes.append({'field': [section], 'before': before, 'after': after})
    if base and list(base['operations']) != list(content['operations']):
        changes.append({'field': ['operation_order'], 'before': list(base['operations']),
                        'after': list(content['operations'])})
    return changes


class PlanStore:
    def __init__(self, root, project, stage):
        self.root = Path(root).resolve()
        self.project, self.stage = identifier(project), identifier(stage)
        self.folder = self.root / 'plans' / self.project / self.stage
        if not self.folder.resolve().is_relative_to(self.root):
            raise PlanConflict('方案目录必须位于工作区内')
        self.path = self.folder / 'plan.json'
        self.plan_id = f'{self.project}:{self.stage}'

    @contextmanager
    def locked(self):
        self.folder.mkdir(parents=True, exist_ok=True)
        lock = self.folder / '.write.lock'
        try:
            fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError as exc:
            raise PlanConflict(f'方案正在写入；若上次异常退出，确认无写入进程后再移除锁：{lock}') from exc
        try:
            with os.fdopen(fd, 'w', encoding='utf-8') as stream:
                stream.write(f'pid={os.getpid()} time={now()}')
            yield
        finally:
            lock.unlink()

    def load(self):
        if not self.path.exists():
            raise PlanConflict(f'尚无方案：{self.plan_id}；请从指定资料创建初始版本')
        data = read_json(self.path)
        if (not isinstance(data, dict) or data.get('schema_version') != SCHEMA or
                data.get('plan_id') != self.plan_id or data.get('project_id') != self.project or
                data.get('stage_id') != self.stage):
            raise PlanConflict('方案主文件版本或身份不一致')
        return data

    def write(self, data):
        atomic_write(self.path, json_bytes(data))

    def resolve(self, data, version):
        if version == 'best':
            version = data['best_version']
            if version is None:
                raise PlanConflict('当前工序尚未指定最优方案；请指定一个已有版本')
        number = version_number(version)
        if str(number) not in data['versions']:
            raise PlanConflict(f'方案版本不存在：{number}')
        return number

    def show(self, version='best'):
        data = self.load()
        number = self.resolve(data, version)
        return dict(plan_id=self.plan_id, project_id=self.project, stage_id=self.stage,
                    stage_name=data['stage_name'], best_version=data['best_version'],
                    **deepcopy(data['versions'][str(number)]),
                    record_links=[r for r in data['record_links'] if r['plan_version'] == str(number)])

    def history(self):
        data = self.load()
        return {k: deepcopy(v) for k, v in data.items() if k not in ('exports', 'operations', 'versions')} | {
            'versions': [{k: v for k, v in version.items() if k != 'content'} for version in data['versions'].values()]}

    def preview(self, command, request):
        if command not in ('create', 'derive'):
            raise PlanConflict('仅预览 create/derive')
        data = self.load() if self.path.exists() else None
        if command == 'create' and data:
            raise PlanConflict('工序方案已存在')
        base_number = None if command == 'create' else self.resolve(data, request.get('base_version'))
        base = {} if base_number is None else data['versions'][str(base_number)]['content']
        content = (validate_content(request.get('content')) if command == 'create'
                   else merge_content(base, request.get('patch'), request.get('operation_order')))
        reviewed_request = {k: v for k, v in request.items() if k not in ('review_digest', 'confirmation')}
        signature = digest(json.dumps([data, command, reviewed_request], ensure_ascii=False, sort_keys=True).encode())
        return {'plan_id': self.plan_id, 'base_version': base_number, 'changes': changes_between(base, content),
                'source': request.get('source', {}), 'review_digest': signature, 'content': content}

    def apply(self, command, request):
        common = {'project_id', 'stage_id', 'operation_id', 'date', 'review_digest', 'confirmation'}
        options = {'create': {'stage_name', 'content', 'mode', 'source'},
                   'derive': {'base_version', 'patch', 'mode', 'source', 'operation_order'},
                   'set-best': {'version', 'instruction'}}
        if command not in options or not isinstance(request, dict) or set(request) - common - options[command]:
            raise PlanConflict('未知操作或输入字段')
        if request.get('project_id') != self.project or request.get('stage_id') != self.stage:
            raise PlanConflict('输入项目和工序不一致')
        operation = nonempty(request.get('operation_id'), '稳定操作标识')
        request_hash = digest(json.dumps([command, request], ensure_ascii=False, sort_keys=True).encode())
        day = request.get('date', date.today().isoformat())
        date.fromisoformat(day)
        with self.locked():
            data = self.load() if self.path.exists() else None
            if data and operation in data['operations']:
                old = data['operations'][operation]
                if old['request_hash'] != request_hash:
                    raise PlanConflict('同一操作标识对应了不同内容；新操作须使用新标识')
                return deepcopy(old['result'])
            if command == 'create':
                if data:
                    raise PlanConflict('工序方案已存在；请基于已有版本修改')
                stage_name = nonempty(request.get('stage_name'), '工序名称')
                for sibling in self.folder.parent.glob('*/plan.json'):
                    other = read_json(sibling)
                    if other['stage_name'].strip().casefold() == stage_name.strip().casefold():
                        raise PlanConflict(f'该工序已登记为 {other["stage_id"]}，请复用已有编号')
                data = dict(schema_version=SCHEMA, project_id=self.project, stage_id=self.stage,
                            stage_name=stage_name, plan_id=self.plan_id, revision=0,
                            best_version=None, best_history=[], versions={}, events=[],
                            record_links=[], operations={}, exports={})
            elif data is None:
                raise PlanConflict('须先创建初始方案')
            if command == 'set-best':
                selected = self.resolve(data, version_number(request.get('version')))
                instruction = nonempty(request.get('instruction'), '用户指定最优的原话')
                previous = data['best_version']
                if previous != selected:
                    data['best_history'].append(dict(previous=previous, version=selected, date=day,
                                                     selected_at=now(), instruction=instruction,
                                                     operation_id=operation))
                    data['best_version'] = selected
                result = dict(plan_id=self.plan_id, version=selected, previous_best=previous,
                              best_version=selected, date=day, created=False)
            else:
                mode = request.get('mode', 'planned')
                if mode not in ('planned', 'observed'):
                    raise PlanConflict('mode 只能为 planned 或 observed')
                source = request.get('source', {})
                if not isinstance(source, dict) or any(not isinstance(v, str) for v in source.values()):
                    raise PlanConflict('source 须为文本字段映射')
                if set(source) - SOURCE_FIELDS:
                    raise PlanConflict('未知来源字段')
                if source.get('source_type') == 'literature':
                    if not all(source.get(k) for k in ('title', 'location', 'locator', 'original_method', 'adaptation_reason')):
                        raise PlanConflict('文献方案须有标题、位置、章节页码、原始方法及适配理由')
                    preview = self.preview(command, request)
                    if request.get('review_digest') != preview['review_digest'] or not request.get('confirmation', '').strip():
                        raise PlanConflict('文献方案须先 preview 展示差异，再提交用户确认原话和 review_digest')
                if mode == 'observed' and not source.get('record_id'):
                    raise PlanConflict('实际条件变化须关联稳定的日报 record_id；日报事实另行保存')
                base_number = None if command == 'create' else self.resolve(data, request.get('base_version'))
                base = {} if base_number is None else data['versions'][str(base_number)]['content']
                content = (validate_content(request.get('content')) if command == 'create'
                           else merge_content(base, request.get('patch'), request.get('operation_order')))
                changes = changes_between(base, content)
                # Identical full conditions reuse a snapshot; each attempt still has its own event/daily link.
                existing = next((int(k) for k, v in data['versions'].items()
                                 if validate_content(v['content']) == content and
                                 list(v['content']['operations']) == list(content['operations'])), None)
                selected = existing or len(data['versions']) + 1
                if existing is None:
                    data['versions'][str(selected)] = dict(version=selected, base_version=base_number,
                        date=day, created_at=now(), mode=mode, source=deepcopy(source),
                        changes=changes, content=content)
                data['events'].append(dict(operation_id=operation, date=day, recorded_at=now(),
                    version=selected, base_version=base_number, mode=mode, source=deepcopy(source),
                    changes=changes, created=existing is None,
                    **({'confirmation': request['confirmation'], 'review_digest': request['review_digest']}
                       if source.get('source_type') == 'literature' else {})))
                result = dict(plan_id=self.plan_id, version=selected, base_version=base_number,
                              best_version=data['best_version'], date=day, created=existing is None)
            data['operations'][operation] = dict(request_hash=request_hash, result=result)
            data['revision'] += 1
            self.write(data)
            return deepcopy(result)

    def export(self, reviewed=None):
        reviewed = reviewed or {}
        if not isinstance(reviewed, dict):
            raise PlanConflict('复核信息须为相对文件名到 SHA256 的映射')
        with self.locked():
            data = self.load()
            names = ['方案总览.docx', *[f'versions/v{int(v):04d}.md' for v in data['versions']]]
            published = {int(Path(name).stem[1:]) for name in data['exports']
                         if re.fullmatch(r'sop/v\d+\.docx', name)}
            if data['best_version'] is not None:
                names.append('SOP.docx')
                if not sop_pending(data['versions'][str(data['best_version'])]['content']):
                    published.add(data['best_version'])
            names.extend(f'sop/v{v:04d}.docx' for v in sorted(published))
            if set(reviewed) - set(names):
                raise PlanConflict('复核清单含未知导出文件')
            checked_hashes = {}
            # Preflight the whole batch before replacing any document.
            for name in names:
                target = self.export_path(name)
                if target.exists():
                    actual = digest(target.read_bytes())
                    checked_hashes[name] = actual
                    known = data['exports'].get(name, {})
                    if actual not in (known.get('sha256'), known.get('pending_sha256')):
                        if reviewed.get(name) != actual:
                            raise PlanConflict(f'Word 存在手动修改或来源未知：{target}；先读取协调，复核 SHA256={actual}')
                        backup = self.export_path(f'manual_edits/{actual}{Path(name).suffix}')
                        if not backup.exists():
                            atomic_write(backup, target.read_bytes())
            written = []
            for name in names:
                target = self.export_path(name)
                if name.endswith('.md'):
                    payload = build_condition_note(data, int(Path(name).stem[1:])).encode('utf-8')
                elif name == '方案总览.docx':
                    payload = build_docx(data)
                else:
                    number = data['best_version'] if name == 'SOP.docx' else int(Path(name).stem[1:])
                    payload = build_docx(data, number, sop=True,
                                         link_prefix='../../../' if name == 'SOP.docx' else '../../../../')
                new_hash = digest(payload)
                current_hash = digest(target.read_bytes()) if target.exists() else None
                if current_hash != checked_hashes.get(name):
                    raise PlanConflict(f'导出检查后文件又发生变化：{target}；请重新读取协调')
                # Journal both hashes so an interrupted replace can be retried safely.
                data['exports'][name] = dict(sha256=current_hash, pending_sha256=new_hash)
                self.write(data)
                atomic_write(target, payload)
                data['exports'][name] = dict(sha256=new_hash, revision=data['revision'])
                self.write(data)
                written.append(str(target))
            return written

    def export_path(self, relative):
        path = self.folder / relative
        if not path.resolve().is_relative_to(self.folder.resolve()):
            raise PlanConflict('导出文件必须位于该工序目录内')
        return path


def validate_record_plan(root, record):
    values = [record.get(key, '') for key in PLAN_FIELDS]
    if not any(values):
        return None
    if not all(values):
        raise PlanConflict('日报方案引用须同时提供 plan_stage、plan_id 和 plan_version')
    if not isinstance(values[1], str) or values[1].count(':') != 1:
        raise PlanConflict('方案编号须使用工具返回的 project:stage')
    project, stage = values[1].split(':')
    if stage != values[0] or project != record['project'].split()[0]:
        raise PlanConflict('日报项目、工序与方案引用不一致')
    store = PlanStore(root, project, stage)
    store.show(version_number(values[2]))
    return store


def sync_record_links(root, records):
    """Rebuild reverse links from authoritative daily metadata; never manufacture execution facts."""
    groups = {}
    for record in records:
        store = validate_record_plan(root, record)
        if store:
            group = groups.setdefault((store.project, store.stage), {})
            group[record['record_id']] = {k: record.get(k, '') for k in (
                'record_id', 'date', 'sample_id', 'step', 'path', *PLAN_FIELDS)}
    # Include formerly linked streams so corrected or removed daily references do not leave stale links.
    for path in (Path(root) / 'plans').glob('*/*/plan.json'):
        data = read_json(path)
        if data.get('record_links'):
            groups.setdefault((data['project_id'], data['stage_id']), {})
    for (project, stage), links in groups.items():
        store = PlanStore(root, project, stage)
        wanted = sorted(links.values(), key=lambda x: (x['date'], x['record_id']))
        with store.locked():
            data = store.load()
            if data['record_links'] != wanted:
                data['record_links'] = wanted
                data['revision'] += 1
                store.write(data)


def text_value(value):
    if value is None:
        return '无'
    if isinstance(value, list):
        return '；'.join(value) or '无'
    return str(value)


def document_table(doc, headings, rows, widths):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    table = doc.add_table(rows=1, cols=len(headings))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for idx, row in enumerate([headings, *rows]):
        cells = table.rows[0].cells if idx == 0 else table.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Cm(width)
            cell.text = text_value(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement('w:' + side)
                edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9')
                borders.append(edge)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement('w:' + side)
                item.set(qn('w:w'), '85'); item.set(qn('w:type'), 'dxa'); margins.append(item)
            props.append(margins)
            if idx == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EEF5'); props.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.keep_with_next = idx == 0
                for run in paragraph.runs:
                    run.bold = idx == 0
                    run.font.size = Pt(10.5)
    return table


def add_content(doc, content):
    content = validate_content(content)
    for section in SECTIONS:
        doc.add_heading(LABELS[section], level=2)
        value = content[section]
        if section in ('purpose', 'scope'):
            doc.add_paragraph(value)
        elif section in MAP_SECTIONS and section != 'operations':
            if value:
                document_table(doc, ['项目', '内容'], list(value.items()), [4.3, 13.2])
            else:
                doc.add_paragraph('待补')
        elif section == 'operations':
            for key, body in value.items():
                paragraph = doc.add_paragraph()
                paragraph.add_run(key + '  ').bold = True
                paragraph.add_run(body)
            if not value:
                doc.add_paragraph('待补')
        else:
            pending = [*value, *[LABELS[s] for s in MAP_SECTIONS if not content[s]]]
            if content['purpose'] == '待补':
                pending.append(LABELS['purpose'])
            doc.add_paragraph('；'.join(dict.fromkeys(pending)) if pending else '无已登记待补项')


def sop_pending(content):
    content = validate_content(content)
    missing = list(content['pending'])
    for section in SECTIONS[:-1]:
        value = content[section]
        values = list(value.values()) if isinstance(value, dict) else [value]
        if not value or any('待补' in v or '待确认' in v or '待确定' in v for v in values):
            missing.append(LABELS[section])
    return list(dict.fromkeys(missing))


def build_condition_note(data, number):
    version = data['versions'][str(number)]
    content = validate_content(version['content'])
    lines = [f"# {data['stage_name']} 条件记录 v{number:04d}", '',
             f"项目：{data['project_id']}；日期：{version['date']}；基线：{version['base_version'] or '初始'}",
             '状态：' + ('当前指定最优' if data['best_version'] == number else '普通／历史版本'), '',
             '执行及实测结果以关联日报为准。', '']
    for section in SECTIONS:
        value = content[section]
        if isinstance(value, dict):
            lines.extend(f'- {LABELS[section]} / {key}：{body}' for key, body in value.items())
        elif isinstance(value, list):
            if value:
                lines.append('- 待补：' + '；'.join(value))
        elif value != '待补':
            lines.append(f'- {LABELS[section]}：{value}')
    for event in data['events']:
        if event['version'] == number and event.get('source'):
            lines.append('- 来源／调整理由：' + json.dumps(event['source'], ensure_ascii=False))
    for record in data['record_links']:
        if record['plan_version'] == str(number):
            from urllib.parse import quote
            lines.append(f"- 关联实验：{record['sample_id']} {record['step']} {record['date']} "
                         f"[日报](../../../../{quote(record['path'])})，记录 {record['record_id']}")
    return '\n'.join(lines) + '\n'


def add_file_link(paragraph, label, target):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    relation = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'), relation)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle'); style.set(qn('w:val'), 'Hyperlink'); props.append(style)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text); hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_docx(data, number=None, sop=False, link_prefix=None):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.59), Cm(27.94)
    section.left_margin = section.right_margin = Cm(2.045)
    section.top_margin = section.bottom_margin = Cm(1.8)
    for name, size in [('Normal', 11), ('Title', 20), ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.paragraph_format.space_after = Pt(4)
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
    selected = data['best_version'] if number is None else number
    links = [r for r in data['record_links'] if number is None or r['plan_version'] == str(number)]
    missing = sop_pending(data['versions'][str(number)]['content']) if sop else []
    title = data['stage_name'] + ('方案总览' if number is None else
                                f' SOP{"待补" if missing else ""} 第{number}版')
    doc.add_paragraph(title, 'Title')
    if number is None:
        doc.add_paragraph(f'项目 {data["project_id"]}    工序 {data["stage_name"]}')
        if selected is None:
            doc.add_paragraph('当前尚未指定最优方案。已有方案见下方版本记录。')
        else:
            doc.add_paragraph(f'当前最优方案为第 {selected} 版，由用户指定。其他尝试见条件记录。')
            paragraph = doc.add_paragraph()
            add_file_link(paragraph, '打开当前 SOP' if not sop_pending(data['versions'][str(selected)]['content'])
                          else '打开最优方案待补稿', 'SOP.docx')
    else:
        version = data['versions'][str(number)]
        mode = '计划采用' if version['mode'] == 'planned' else '来自实际条件变更记录'
        doc.add_paragraph(f'项目 {data["project_id"]}    日期 {version["date"]}    {mode}')
        doc.add_paragraph('方案状态：' + ('当前指定最优' if number == data['best_version'] else '历史 SOP，已被替换'))
        designation = next((item for item in reversed(data['best_history']) if item['version'] == number), None)
        if designation:
            doc.add_paragraph(f"用户指定：{designation['date']}；{designation['instruction']}")
        if missing:
            doc.add_paragraph('SOP 待补信息：' + '；'.join(missing) + '。补齐前不作为正式 SOP。')
        if version['base_version'] is not None:
            doc.add_paragraph(f'基于第 {version["base_version"]} 版调整，未修改项目继承自该版。')
        doc.add_paragraph('配方和条件用于执行参考；本版是否已执行、实际用量及结果以关联日报为准。')
    if selected is not None and sop:
        add_content(doc, data['versions'][str(selected)]['content'])
    change_heading = doc.add_heading('变更记录', level=1 if number is None else 2)
    if number is None:
        for event in data['events']:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            mode = '计划' if event['mode'] == 'planned' else '实际条件记录'
            base = f'基于第 {event["base_version"]} 版' if event['base_version'] else '初始方案'
            paragraph.add_run(f'{event["date"]}  {mode}  {base}  ').bold = True
            add_file_link(paragraph, f'打开第 {event["version"]} 版条件记录', f'versions/v{event["version"]:04d}.md')
            if event.get('source'):
                source = event['source']
                doc.add_paragraph('来源／调整理由：' + '；'.join(source.values()))
            changed = event['changes'] if event['base_version'] else []
            if changed:
                document_table(doc, ['变更项目', '原条件', '调整后'], [
                    [' / '.join([LABELS[c['field'][0]], *c['field'][1:]]), c['before'], c['after']]
                    for c in changed], [5.0, 6.25, 6.25])
            else:
                doc.add_paragraph('建立初始方案。' if event['base_version'] is None else '条件未变，复用已有版本。')
    else:
        version = data['versions'][str(number)]
        if version['base_version'] is None:
            doc.add_paragraph('初始方案。')
        else:
            for change in version['changes']:
                field = ' / '.join([LABELS[change['field'][0]], *change['field'][1:]])
                doc.add_paragraph(f'{field}：{text_value(change["before"])} → {text_value(change["after"])}')
        source = dict(version['source'])
        if any(r['record_id'] == source.get('record_id') and
               (not source.get('sample_id') or r['sample_id'] == source['sample_id']) for r in links):
            source.pop('record_id', None)
            source.pop('sample_id', None)
        if source:
            source_labels = {'description': '说明', 'location': '资料位置', 'record_id': '日报记录',
                             'sample_id': '样品', 'context_version': '背景版本', 'source_type': '来源类型',
                             'title': '文献', 'locator': '页码／章节', 'original_method': '文献原始方法',
                             'adaptation_reason': '适配理由'}
            doc.add_paragraph('来源：' + '；'.join(f'{source_labels[key]} {value}' for key, value in source.items()))
    if number is None and data['best_history']:
        doc.add_heading('最优标记记录', level=1)
        document_table(doc, ['日期', '指定版本', '用户说明'], [
            [item['date'], str(item['version']), item['instruction']] for item in data['best_history']], [3, 2.5, 12])
    if links:
        if number is None:
            doc.add_heading('关联实验记录', level=1)
        for record in links:
            lead = '' if number is None else '关联实验：'
            paragraph = doc.add_paragraph(f'{lead}{record["date"]}  {record["sample_id"]}  {record["step"]}  第 {record["plan_version"]} 版  ')
            prefix = link_prefix or ('../../../' if number is None else '../../../../')
            from urllib.parse import quote
            add_file_link(paragraph, '打开日报', prefix + quote(record['path'].replace('\\', '/')))
    buffer = io.BytesIO(); doc.save(buffer)
    return buffer.getvalue()


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=None)
    sub = parser.add_subparsers(dest='command', required=True)
    for command in ('create', 'derive', 'set-best'):
        sub.add_parser(command).add_argument('input', type=Path)
    child = sub.add_parser('preview')
    child.add_argument('input', type=Path)
    child.add_argument('--action', choices=['create', 'derive'], required=True)
    for command in ('show', 'history', 'export'):
        child = sub.add_parser(command)
        child.add_argument('--project', required=True); child.add_argument('--stage', required=True)
        if command == 'show':
            child.add_argument('--version', default='best')
        if command == 'export':
            child.add_argument('--reviewed', type=Path, help='Reviewed filename/SHA256 map, only after inspecting and reconciling manual edits')
    child = sub.add_parser('list'); child.add_argument('--project', required=True)
    args = parser.parse_args()
    args.root = data_root(args.root)
    try:
        if args.command in ('create', 'derive', 'set-best', 'preview'):
            request = read_json(args.input)
            if not isinstance(request, dict):
                raise PlanConflict('输入须为 JSON 对象')
            store = PlanStore(args.root, request.get('project_id'), request.get('stage_id'))
            output = store.preview(args.action, request) if args.command == 'preview' else store.apply(args.command, request)
            if args.command != 'preview':
                from project_context import route_status
                output['route_check'] = route_status(args.root, store.project)
        elif args.command == 'list':
            project = identifier(args.project)
            directory = (args.root / 'plans' / project).resolve()
            if not directory.is_relative_to(args.root.resolve()):
                raise PlanConflict('项目目录必须位于工作区内')
            output = []
            for path in sorted(directory.glob('*/plan.json')):
                data = read_json(path)
                output.append({k: data[k] for k in ('plan_id', 'stage_id', 'stage_name', 'best_version')})
        else:
            store = PlanStore(args.root, args.project, args.stage)
            if args.command == 'show':
                output = store.show(args.version)
            elif args.command == 'history':
                output = store.history()
            else:
                output = {'paths': store.export(read_json(args.reviewed) if args.reviewed else None)}
        print(json.dumps(output, ensure_ascii=False, indent=2))
    except (ValueError, OSError, KeyError, TypeError) as exc:
        print(json.dumps({'error': str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
