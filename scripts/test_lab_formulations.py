"""Fictional arithmetic and batch lifecycle checks; never use real research data."""
from copy import deepcopy
from decimal import Decimal
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from lab_formulation_calc import calculate, expand
from lab_formulations import BatchStore, today
from lab_plans import PlanStore
from lab_records import Store


def q(value, unit='g'):
    return dict(value=str(value), unit=unit, source='软件测试输入')


def materials():
    return [dict(id=x, name='测试材料' + x, material_type='pure', properties={}) for x in ('A', 'B', 'S')]


def recipe(calcs, mats=None):
    return dict(materials=mats or materials(), calculations=calcs)


def values(result):
    return {mid: Decimal(result['nodes'][out['node']]['value']) for mid, out in result['outputs'].items()}


def spec():
    return dict(project_id='DEMO', stage_id='mix', stage_name='软件测试混合', batch_id='trial',
        operation_id='save-1', description='虚构软件测试：质量比1:2，总量30g，各组重复两次。', date='2030-01-07',
        recipe=recipe([dict(mode='mass_ratio', total=q(30), ratios={'A': '1', 'B': '2'})], materials()[:2]),
        repeats=2, plan_content=dict(purpose='软件测试', tests={'测试': '方法待补'}))


class ArithmeticTest(unittest.TestCase):
    def calc(self, calcs, mats=None):
        return calculate(recipe(calcs, mats), '软件测试，非实验条件')

    def test_ratios_zero_and_units(self):
        result = self.calc([dict(mode='mass_ratio', total=q(30000, 'mg'), ratios={'A': 1, 'B': 2, 'S': 0})])
        self.assertEqual(values(result), {'A': 10, 'B': 20, 'S': 0})
        with self.assertRaises(ValueError):
            self.calc([dict(mode='mass_ratio', total=q(30), ratios={'A': 0, 'B': 0})])

    def test_fractions_and_make_up(self):
        result = self.calc([dict(mode='mass_fraction', total=q(100), fractions={'A': q(20, '%'), 'B': q('.3', 'fraction')}),
            dict(mode='make_up', total=q(100), material='S', members=['A', 'B'])])
        self.assertEqual(values(result), {'A': 20, 'B': 30, 'S': 50})
        with self.assertRaises(ValueError):
            self.calc([dict(mode='mass_fraction', total=q(10), fractions={'A': q(101, '%')})])

    def test_solids_and_effective_component(self):
        mats = materials()
        mats[0]['properties']['solids'] = q(50, '%')
        mats[1]['properties']['solids'] = q(100, '%')
        mats[2]['properties']['solids'] = q(0, '%')
        result = self.calc([dict(mode='solids', total=q(100), fraction=q(30, '%'), ratios={'A': 1, 'B': 2}, solvent='S')], mats)
        self.assertEqual(values(result), {'A': 20, 'B': 20, 'S': 60})
        result = self.calc([dict(mode='mass_fraction', total=q(100), fractions={'A': q(10, '%')},
                                basis='component', content_property='solids')], mats)
        self.assertEqual(values(result)['A'], 20)
        with self.assertRaises(ValueError):
            self.calc([dict(mode='solids', total=q(10), fraction=q(90, '%'), ratios={'A': 1}, solvent='S')], mats)

    def test_missing_properties_preserves_known_outputs(self):
        result = self.calc([dict(mode='fixed', material='A', quantity=q(5)), dict(mode='molar_ratio', reference='A', ratios={'A': 1, 'B': 2})])
        self.assertEqual(values(result), {'A': 5})
        self.assertTrue(any('molar_mass' in s for s in result['pending']))

    def test_dry_total_and_component_ratio(self):
        mats = materials()
        mats[0]['properties']['solids'] = q(50, '%')
        mats[1]['properties']['solids'] = q(100, '%')
        mats[2]['properties']['solids'] = q(0, '%')
        result = self.calc([dict(mode='solids', total=q(30), total_basis='dry', fraction=q(30, '%'), ratios={'A': 1, 'B': 2}, solvent='S')], mats)
        self.assertEqual(values(result), {'A': 20, 'B': 20, 'S': 60})
        result = self.calc([dict(mode='mass_ratio', total=q(30), ratios={'A': 1, 'B': 2}, basis='component', content_property='solids')], mats)
        self.assertEqual(values(result), {'A': 20, 'B': 20})

    def test_zero_molar_addition_and_material_equivalent_basis(self):
        mats = materials()[:2]
        mats[0]['properties']['molar_mass'] = q(100, 'g/mol')
        result = self.calc([dict(mode='fixed', material='A', quantity=q(1)), dict(mode='molar_ratio', reference='A', ratios={'A': 1, 'B': 0})], mats)
        self.assertEqual(values(result)['B'], 0)
        mats[0].update(material_type='mixture', properties={'equivalent_weight': q(100, 'g/eq')})
        mats[1].update(material_type='mixture', properties={'equivalent_weight': q(50, 'g/eq')})
        result = self.calc([dict(mode='fixed', material='A', quantity=q(2)), dict(mode='equivalents', reference='A', ratios={'A': 1, 'B': 1})], mats)
        self.assertEqual(values(result), {'A': 2, 'B': 1})

    def test_invalid_properties_and_duplicate_assignments(self):
        mats = materials()
        mats[0]['properties']['solids'] = q(150, '%')
        with self.assertRaises(ValueError):
            self.calc([dict(mode='fixed', material='A', quantity=q(1))], mats)
        with self.assertRaises(ValueError):
            self.calc([dict(mode='fixed', material='A', quantity=q(1)), dict(mode='fixed', material='A', quantity=q(2))])

    def test_moles_and_purity(self):
        mats = materials()[:2]
        mats[0]['properties'] = dict(molar_mass=q(100, 'g/mol'), purity=q(50, '%'))
        mats[1]['properties'] = dict(molar_mass=q(200, 'g/mol'), purity=q(80, '%'))
        result = self.calc([dict(mode='fixed', material='A', quantity=q(10)), dict(mode='molar_ratio', reference='A', ratios={'A': 1, 'B': 2})], mats)
        self.assertEqual(values(result), {'A': 10, 'B': 25})

    def test_equivalents_by_functionality_and_equivalent_weight(self):
        mats = materials()[:2]
        mats[0]['properties'] = dict(molar_mass=q(100, 'g/mol'), functionality=q(2, '1'))
        mats[1]['properties'] = dict(equivalent_weight=q(50, 'g/eq'))
        result = self.calc([dict(mode='fixed', material='A', quantity=q(1)), dict(mode='equivalents', reference='A', ratios={'A': 1, 'B': 1})], mats)
        self.assertEqual(values(result), {'A': 1, 'B': 1})
        mats[0]['properties'] = dict(equivalent_weight=q(25, 'g/eq'))
        self.assertEqual(values(self.calc([dict(mode='fixed', material='A', quantity=q(1)),
            dict(mode='equivalents', reference='A', ratios={'A': 1, 'B': 1})], mats))['B'], 2)

    def test_solution_molar_basis(self):
        mats = materials()[:2]
        mats[0]['properties'] = dict(molar_mass=q(100, 'g/mol'))
        mats[1].update(material_type='solution', component='虚构组分', properties=dict(molarity=q(1, 'mol/L')))
        result = self.calc([dict(mode='fixed', material='A', quantity=q(1)), dict(mode='molar_ratio', reference='A', ratios={'A': 1, 'B': 2})], mats)
        self.assertEqual(values(result)['B'], 20)
        self.assertEqual(result['outputs']['B']['kind'], 'volume')

    def test_dilution_and_scale(self):
        mats = materials()[:1]
        mats[0].update(material_type='solution', component='虚构组分')
        result = self.calc([dict(mode='dilution', material='A', stock=q(1, 'mol/L'), target=q(100, 'mmol/L'), final_volume=q(.1, 'L'))], mats)
        self.assertEqual(values(result)['A'], 10)
        self.assertTrue(any('定容' in s for s in result['pending']))
        result = self.calc([dict(mode='scale', quantities={'A': q(100, 'mg'), 'B': q(2, 'mL')}, factor=q(3, '1'))])
        self.assertEqual(values(result), {'A': Decimal('.3'), 'B': 6})
        mats = materials()
        mats[0]['properties']['density'] = q(2, 'g/mL')
        result = self.calc([dict(mode='scale', quantities={'A': q(1, 'mL'), 'B': q(3)}, target_total=q(20))], mats)
        self.assertEqual(values(result), {'A': 4, 'B': 12})

    def test_negative_make_up_and_dimension_conflict(self):
        with self.assertRaises(ValueError):
            self.calc([dict(mode='fixed', material='A', quantity=q(20)), dict(mode='make_up', total=q(10), material='S', members=['A'])])
        with self.assertRaises(ValueError):
            self.calc([dict(mode='mass_ratio', total=q(10, 'mL'), ratios={'A': 1})])
        with self.assertRaises(ValueError):
            self.calc([dict(mode='fixed', material='A', quantity=q(-1))])

    def test_expansion_zip_product_and_repetition(self):
        request = spec()
        request['variables'] = [dict(name='A比', path=['calculations', 0, 'ratios', 'A'], values=[0, 1, 2]),
                                dict(name='B比', path=['calculations', 0, 'ratios', 'B'], values=[1, 2, 3])]
        with self.assertRaises(ValueError):
            expand(request)
        request['combination'] = 'zip'
        self.assertEqual(len(expand(request)), 6)
        request['combination'] = 'product'
        self.assertEqual(len(expand(request)), 18)
        self.assertEqual(request['recipe']['calculations'][0]['ratios']['A'], '1')


class BatchTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.store = BatchStore(self.root, 'DEMO', 'mix', 'trial')

    def tearDown(self):
        self.tmp.cleanup()

    def test_preview_has_no_writes_and_save_has_no_experiments(self):
        before = list(self.root.rglob('*'))
        self.store.preview(spec())
        self.assertEqual(list(self.root.rglob('*')), before)
        saved = self.store.save(spec())
        self.assertEqual(saved['groups'][0]['plan_version'], saved['groups'][1]['plan_version'])
        self.assertFalse((self.root / 'report').exists())
        self.assertIsNone(self.store.plans.load()['best_version'])
        self.assertEqual(self.store.save(spec()), saved)

    def test_versions_and_baseline(self):
        self.store.save(spec())
        with self.assertRaises(ValueError):
            self.store.show(0)
        new = spec()
        new.update(operation_id='save-2', base_version=1)
        new['recipe']['calculations'][0]['total'] = q(60)
        self.store.save(new)
        self.assertEqual(values(self.store.show(1)['groups'][0]['result'])['A'], 10)
        self.assertEqual(values(self.store.show(2)['groups'][0]['result'])['A'], 20)
        new['operation_id'] = 'save-3'
        new.pop('base_version')
        with self.assertRaises(ValueError):
            self.store.save(new)

    def test_partial_completion_retry_and_cold_rebuild(self):
        self.store.save(spec())
        req = dict(operation_id='complete-1', confirmation='第一组已完成', groups=['G001-R01'], defaults={'date': '2030-01-08'})
        first = self.store.complete(req)
        self.assertEqual(self.store.complete(req), first)
        records = Store(self.root).query(start='2030-01-01', end='2030-02-28')
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]['state'], '待测试')
        self.assertTrue(all(r['quantity_kind'] == 'planned' for r in records[0]['materials']))
        self.assertIsNone(self.store.show()['completions']['G001-R02'])
        self.store.complete(dict(operation_id='complete-2', confirmation='第二组已完成', groups=['G001-R02'], defaults={'date': '2030-02-01'}))
        (self.root / 'report/index.json').unlink()
        rebuilt = Store(self.root)
        rebuilt.refresh()
        records = rebuilt.query(start='2030-01-01', end='2030-02-28')
        self.assertEqual(len(records), 2)
        self.assertEqual({r['batch_group'] for r in records}, {'G001-R01', 'G001-R02'})
        self.assertEqual({r['batch_revision'] for r in records}, {'1'})

    def test_all_and_repeated_confirmation_does_not_duplicate(self):
        self.store.save(spec())
        req = dict(operation_id='all', confirmation='这批完成', defaults={'date': '2030-01-08'})
        self.store.complete(req)
        self.store.complete(req | dict(operation_id='same-report-new-operation'))
        self.store.complete(dict(operation_id='again', confirmation='这批完成'))
        self.assertEqual(len(Store(self.root).query(start='2030-01-01', end='2030-01-31')), 2)
        self.assertEqual(len({r['sample_id'] for r in Store(self.root).query(start='2030-01-01', end='2030-01-31')}), 2)

    def test_cross_day_and_actual_change(self):
        self.store.save(spec())
        actual = dict(operation='投料', name='测试材料A', role='原料', original='实际称取11g', quantity_kind='actual',
                      material_type='pure', measurements=[dict(kind='mass', value='11', unit='g', source='用户报告')])
        req = dict(operation_id='done', confirmation='两组完成，第一组实际有变化', reports={
            'G001-R01': dict(date='2030-01-08', materials=[actual], as_planned=False),
            'G001-R02': dict(date='2030-01-09', as_planned=True)})
        result = self.store.complete(req)
        self.assertEqual(len(result['paths']), 2)
        rows = Store(self.root).query(start='2030-01-01', end='2030-01-31')
        self.assertNotEqual(rows[0]['plan_version'], rows[1]['plan_version'])
        self.assertEqual([r['quantity_kind'] for r in rows[0]['materials']], ['planned', 'planned', 'actual'])

    def test_daily_interruption_reuses_ids_and_sample_reservation(self):
        self.store.save(spec())
        req = dict(operation_id='done', confirmation='完成', defaults={'date': '2030-01-08'})
        original = Store.save
        def fail_after_word(store, records):
            original(store, records)
            raise OSError('模拟Word已保存后失败')
        with patch.object(Store, 'save', fail_after_word):
            with self.assertRaises(OSError):
                self.store.complete(req)
        before = self.store.load()['operations']['done']['records']
        self.store.complete(req)
        rows = Store(self.root).query(start='2030-01-01', end='2030-01-31')
        self.assertEqual({r['record_id'] for r in rows}, {r['record_id'] for r in before})
        self.assertEqual(len(rows), 2)

    def test_plan_write_interruption_reuses_original_create(self):
        original = PlanStore.apply
        failed = [False]
        def fail_after_create(store, command, request):
            result = original(store, command, request)
            if not failed[0]:
                failed[0] = True
                raise OSError('模拟方案保存后响应失败')
            return result
        with patch.object(PlanStore, 'apply', fail_after_create):
            with self.assertRaises(OSError):
                self.store.save(spec())
        self.store.save(spec())
        self.assertEqual(len(self.store.plans.load()['versions']), 1)

    def test_manual_daily_content_is_preserved(self):
        self.store.save(spec())
        self.store.complete(dict(operation_id='first', confirmation='完成', groups=['G001-R01'], defaults={'date': '2030-01-08'}))
        path = self.root / 'report/2030-01/日报_2030-01-08.docx'
        doc = Document(path)
        doc.add_paragraph('手动补充，必须保留')
        doc.save(path)
        with self.assertRaises(ValueError):
            self.store.complete(dict(operation_id='second', confirmation='完成', groups=['G001-R02'], defaults={'date': '2030-01-08'}))
        self.assertIn('手动补充，必须保留', '\n'.join(p.text for p in Document(path).paragraphs))

    def test_confirmation_and_operation_conflicts(self):
        self.store.save(spec())
        with self.assertRaises(ValueError):
            self.store.complete(dict(operation_id='done', confirmation=''))
        changed = spec()
        changed['repeats'] = 3
        with self.assertRaises(ValueError):
            self.store.save(changed)
        self.assertRegex(today('+08:00'), r'^\d{4}-\d{2}-\d{2}$')
        with self.assertRaises(ValueError):
            today('+08:90')

    def test_invalid_input_does_not_leave_unrecoverable_journal(self):
        bad = spec()
        bad['plan_content'] = {'conditions': {'温度': 50}}
        with self.assertRaises(ValueError):
            self.store.save(bad)
        self.assertFalse(self.store.path.exists())
        with self.assertRaises(ValueError):
            self.store.save(spec() | dict(stage_name=''))
        self.assertFalse(self.store.path.exists())
        self.store.save(spec())
        for report in ({'experiment_type': '不存在'}, {'as_planned': 'yes'}, {'observed_patch': {'bad': 'field'}}):
            with self.assertRaises(ValueError):
                self.store.complete(dict(operation_id='invalid', confirmation='完成', defaults=report))
            self.assertNotIn('invalid', self.store.load()['operations'])

    def test_cross_batch_sample_number_reservations(self):
        self.store.save(spec())
        req = dict(operation_id='interrupted', confirmation='完成', defaults={'date': '2030-01-08'})
        with patch.object(Store, 'save', side_effect=OSError('模拟写入前失败')):
            with self.assertRaises(OSError):
                self.store.complete(req)
        other = BatchStore(self.root, 'DEMO', 'mix', 'another')
        other_request = spec() | dict(batch_id='another', operation_id='another-save', base_version=1, repeats=1)
        other.save(other_request)
        result = other.complete(dict(operation_id='another-complete', confirmation='完成', defaults={'date': '2030-01-08'}))
        self.assertEqual(result['records'][0]['sample_id'], 'DEMO-20300108-3')
        self.store.complete(req)
        self.assertEqual(len(Store(self.root).query(start='2030-01-08', end='2030-01-08')), 3)

    def test_list_ignores_excel_export_sidecars(self):
        self.store.save(spec())
        (self.store.folder / 'trial_v0001.export.json').write_text('{}', encoding='utf-8')
        (self.store.folder / 'trial_v0001.qa.json').write_text('{}', encoding='utf-8')
        import subprocess
        import sys
        output = subprocess.run([sys.executable, str(Path(__file__).with_name('lab_formulations.py')),
            '--root', str(self.root), 'list', '--project', 'DEMO'], check=True, capture_output=True,
            text=True, encoding='utf-8')
        self.assertEqual([r['batch_id'] for r in json.loads(output.stdout)], ['trial'])

    def test_precise_inputs_are_not_deduplicated_by_display_rounding(self):
        first = spec() | dict(repeats=1)
        first['recipe']['calculations'][0]['total'] = q('30.00000001')
        self.store.save(first)
        second = deepcopy(first) | dict(operation_id='precise-update', base_version=1)
        second['recipe']['calculations'][0]['total'] = q('30.00000002')
        updated = self.store.save(second)
        self.assertNotEqual(self.store.show(1)['groups'][0]['plan_version'], updated['groups'][0]['plan_version'])


if __name__ == '__main__':
    unittest.main()
