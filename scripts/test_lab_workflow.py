"""Quantity, confirmation, route and SOP scenarios isolated from real lab data."""
from copy import deepcopy
from decimal import Decimal
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from extract_daily_docx import extract_docx
from lab_quantities import prepare_materials, raw_materials, quantities_hash
from lab_records import Store, Conflict
from lab_plans import PlanStore, PlanConflict
from project_context import update, preview, route_status
from test_lab_records import record
from test_lab_plans import initial, request, create_plan, daily
import yaml


def prop(value, unit):
    return {'value': str(value), 'unit': unit, 'source': '软件测试常数，非化学资料'}


def material(kind='pure', quantity='mass', value='1.00', unit='g', properties=None, **extra):
    row = {'operation': '投料', 'name': '测试材料', 'role': '原料', 'original': value + ' ' + unit,
           'material_type': kind, 'quantity_kind': 'actual',
           'measurements': [{'kind': quantity, 'value': value, 'unit': unit, 'source': '用户原始报告'}],
           'properties': properties or {}}
    row.update(extra)
    return row


def calculated(row, kind, basis='material'):
    return next((Decimal(q['value']) for q in row['calculated'] if q['kind'] == kind and q['basis'] == basis), None)


def complete():
    content = initial()
    content.update(scope='仅用于软件验证的条件示例。', equipment={'装置': '测试装置 A'},
        controls={'记录': '检查温度设定与记录一致。'}, acceptance={'外观': '按确认的示例判定表登记结果。'},
        pending=[])
    content['operations']['3 后处理'] = '按确认的示例条件结束并记录。'
    content['tests']['羟基含量'] = '使用确认的示例方法，记录 mmol/g。'
    return content


class QuantitiesTest(unittest.TestCase):
    def test_solid_nominal_purity_and_liquid_forward_reverse(self):
        row = material(properties={'molar_mass': prop(100, 'g/mol'), 'purity': prop(80, '%')})
        saved = deepcopy(row)
        parsed = prepare_materials([row])[0]
        self.assertEqual(calculated(parsed, 'amount'), 8)
        self.assertEqual(row, saved)
        self.assertEqual(parsed['measurements'][0]['value'], '1.00')
        liquid = prepare_materials([material(quantity='volume', value='2.00', unit='mL',
            properties={'density': prop('.8', 'g/mL'), 'molar_mass': prop(80, 'g/mol')})])[0]
        self.assertEqual(calculated(liquid, 'mass'), Decimal('1.6'))
        self.assertEqual(calculated(liquid, 'amount'), 20)
        reverse = prepare_materials([material(quantity='amount', value='20', unit='mmol',
            properties={'density': prop('.8', 'g/mL'), 'molar_mass': prop(80, 'g/mol')})])[0]
        self.assertEqual(calculated(reverse, 'volume'), 2)

    def test_stock_solution_total_and_solute_mass_do_not_mix(self):
        row = prepare_materials([material(kind='solution', quantity='volume', value='10.0', unit='mL',
            component='测试溶质', properties={'molarity': prop('.1', 'mol/L'),
                'density': prop(1, 'g/mL'), 'molar_mass': prop(100, 'g/mol')})])[0]
        self.assertEqual(calculated(row, 'amount', 'component'), 1)
        self.assertEqual(calculated(row, 'mass', 'component'), Decimal('.1'))
        self.assertEqual(calculated(row, 'mass', 'material'), 10)
        self.assertIsNone(calculated(row, 'amount', 'material'))

    def test_dispersion_and_product_missing_basis(self):
        row = prepare_materials([material(kind='dispersion', quantity='volume', value='10', unit='mL',
            component='测试颗粒', properties={'mass_concentration': prop(50, 'mg/mL'),
                                           'molar_mass': prop(100, 'g/mol')})])[0]
        self.assertEqual(calculated(row, 'mass', 'component'), Decimal('.5'))
        self.assertIsNone(calculated(row, 'amount', 'component'))
        unknown = prepare_materials([material(kind='unknown', role='产物', properties={'molar_mass': prop(100, 'g/mol')})])[0]
        self.assertFalse(unknown['calculated'])
        self.assertTrue(unknown['pending'])

    def test_conflicting_measurements_withhold_calculation_and_reject_units(self):
        row = material(properties={'molar_mass': prop(100, 'g/mol')})
        row['measurements'].append({'kind': 'amount', 'value': '20', 'unit': 'mmol', 'source': '用户另一报告'})
        parsed = prepare_materials([row])[0]
        self.assertTrue(parsed['conflicts'])
        self.assertFalse(parsed['calculated'])
        for invalid in (material(unit='mL'), material(value='NaN'), material(value='-1'),
                        material(properties={'density': {'value': 1, 'unit': 'g/mL'}})):
            with self.assertRaises(ValueError):
                prepare_materials([invalid])


class WorkflowTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_word_roundtrip_cold_rebuild_and_narrative_correction_preserve_table(self):
        row = material(properties={'molar_mass': prop(100, 'g/mol')})
        store = Store(self.root)
        path = store.save([record(materials=[row])])
        source = extract_docx(path)['records'][0]
        self.assertEqual(source['materials'][0]['original'], '1.00 g')
        self.assertEqual(calculated(source['materials'][0], 'amount'), 10)
        store.index_path.unlink()
        cold = Store(self.root); cold.refresh()
        self.assertEqual(cold.query(sample='DEMO-20260917-1')[0]['materials'], source['materials'])
        cold.save([record(today='补充了操作说明。')])
        self.assertEqual(extract_docx(path)['records'][0]['materials'], source['materials'])
        # Appending another card must retain the first card's nested table and data.
        cold.save([record('r2', sample='DEMO-20260917-2')])
        self.assertEqual(extract_docx(path)['records'][0]['materials'], source['materials'])

    def test_record_correction_requires_review_and_manual_table_edits_are_detected(self):
        row = material()
        store = Store(self.root); path = store.save([record(materials=[row])])
        new = material(value='2.00')
        before_bytes = path.read_bytes()
        with self.assertRaisesRegex(Conflict, '用量冲突'):
            store.save([record(materials=[new])])
        self.assertEqual(path.read_bytes(), before_bytes)
        old_rows = extract_docx(path)['records'][0]['materials']
        store.save([record(materials=[new], reviewed_materials=quantities_hash(raw_materials(old_rows)))])
        self.assertEqual(extract_docx(path)['records'][0]['materials'][0]['original'], '2.00 g')
        doc = Document(path)
        doc.tables[0].cell(0, 0).tables[0].cell(1, 2).text = '3.00 g'
        doc.save(path)
        with self.assertRaisesRegex(ValueError, '手动修改'):
            store.refresh()

    def test_branch_sample_and_planned_actual_values_remain_separate(self):
        store = Store(self.root); store.save([record()])
        actual = material(sample_id='DEMO-20260917-1-B01', sample_step='S01')
        planned = material(quantity_kind='planned', value='2.00')
        branch = record('branch', sample='DEMO-20260917-1-B01', source_sample='DEMO-20260917-1',
                        source_step='S01', materials=[actual, planned])
        store.save([branch])
        found = store.query(sample=branch['sample_id'])[0]
        self.assertEqual(found['source_sample'], 'DEMO-20260917-1')
        self.assertEqual([r['quantity_kind'] for r in found['materials']], ['actual', 'planned'])
        self.assertEqual(found['materials'][0]['sample_id'], branch['sample_id'])

    def candidate(self, sources=None):
        data = dict(project_id='DEMO', project_name='软件验证项目', status='active',
                    project_route='待建立', targets={}, chart_groups={}, reference_sources=sources or [])
        path = self.root / 'candidate.yaml'
        path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding='utf-8')
        return path, data

    def test_literature_context_preview_is_readonly_and_confirmation_is_bound(self):
        path, data = self.candidate([dict(source_type='literature', title='测试文献', location='fixture.pdf',
            purpose='软件行为验证', locator='第 2 页', status='available')])
        review = preview(self.root, path)
        self.assertFalse((self.root / 'references').exists())
        with self.assertRaisesRegex(ValueError, 'preview'):
            update(self.root, path, '2026-10-08')
        data['project_route'] = '新的候选路线'
        path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'preview'):
            update(self.root, path, '2026-10-08', review_digest=review['review_digest'], confirmation='确认')
        review = preview(self.root, path)
        result = update(self.root, path, '2026-10-08', review_digest=review['review_digest'], confirmation='确认该文献背景草案')
        saved = yaml.safe_load(result.read_text(encoding='utf-8'))
        self.assertEqual(saved['context_change']['confirmation'], '确认该文献背景草案')
        self.assertEqual(saved['context_version'], 1)

    def test_missing_route_questions_are_deduplicated_without_blocking_facts(self):
        path, data = self.candidate(); update(self.root, path, '2026-10-08')
        check = route_status(self.root, 'DEMO')
        self.assertTrue(check['needs_communication'])
        Store(self.root).save([record()])
        route_status(self.root, 'DEMO', True, check['question_id'])
        self.assertFalse(route_status(self.root, 'DEMO')['needs_communication'])
        data['project_route'] = '目标调整后的路线待确认'
        path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding='utf-8')
        update(self.root, path, '2026-10-08')
        self.assertTrue(route_status(self.root, 'DEMO')['needs_communication'])
        data['technical_route'] = {'status': 'confirmed', 'stages': [{k: '已确认内容' for k in
            ('id', 'name', 'purpose', 'method', 'variables', 'inputs', 'outputs', 'metrics', 'decision', 'next')}]}
        path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding='utf-8')
        update(self.root, path, '2026-10-08')
        self.assertEqual(route_status(self.root, 'DEMO')['status'], 'confirmed')

    def test_literature_plan_confirmation_retry_and_stale_baseline(self):
        plan = create_plan(self.root)
        source = dict(source_type='literature', title='测试文献', location='fixture.pdf', locator='方法第 2 节',
                      original_method='原文温度 70 ℃', adaptation_reason='根据用户设备修改温度')
        payload = request('literature', base_version=1, patch={'conditions': {'反应温度': '70 ℃'}}, source=source)
        old_bytes = plan.path.read_bytes()
        review = plan.preview('derive', payload)
        self.assertEqual(plan.path.read_bytes(), old_bytes)
        with self.assertRaisesRegex(PlanConflict, 'preview'):
            plan.apply('derive', payload)
        payload.update(review_digest=review['review_digest'], confirmation='确认该文献方案调整')
        result = plan.apply('derive', payload)
        self.assertEqual(plan.apply('derive', payload), result)
        stale = request('stale', base_version=1, patch={'conditions': {'反应时间': '5 h'}}, source=source)
        review = plan.preview('derive', stale)
        plan.apply('derive', request('other', base_version=1, patch={'conditions': {'反应时间': '7 h'}}))
        stale.update(review_digest=review['review_digest'], confirmation='确认')
        with self.assertRaisesRegex(PlanConflict, 'preview'):
            plan.apply('derive', stale)

    def test_only_designated_best_publishes_sop_and_switch_preserves_references(self):
        plan = PlanStore(self.root, 'DEMO', 'peg')
        plan.apply('create', request('init', stage_name='示例工序', content=complete()))
        plan.export()
        self.assertFalse((plan.folder / 'SOP.docx').exists())
        self.assertTrue((plan.folder / 'versions/v0001.md').exists())
        self.assertFalse((plan.folder / 'versions/v0001.docx').exists())
        plan.apply('set-best', request('best-1', version=1, instruction='指定第一版最优'))
        Store(self.root).save([daily(version='1')]); plan.export()
        first = plan.show(1)['content']
        plan.apply('derive', request('change', base_version=1, patch={'conditions': {'反应温度': '80 ℃'}}))
        plan.apply('set-best', request('best-2', version=2, instruction='第二版作为最优'))
        plan.export()
        self.assertEqual(plan.show(1)['content'], first)
        self.assertEqual(Store(self.root).query(sample='DEMO-20260920-1')[0]['plan_version'], '1')
        old_sop = Document(plan.folder / 'sop/v0001.docx')
        self.assertIn('历史 SOP', '\n'.join(p.text for p in old_sop.paragraphs))
        current = Document(plan.folder / 'SOP.docx')
        self.assertIn('当前指定最优', '\n'.join(p.text for p in current.paragraphs))
        relationships = [r.target_ref for r in current.part.rels.values() if 'hyperlink' in r.reltype]
        self.assertTrue(any(x.startswith('../../../report/') for x in relationships) or not plan.show(2)['record_links'])

    def test_incomplete_best_is_pending_draft_not_published_sop(self):
        plan = create_plan(self.root)
        plan.apply('set-best', request('best', version=1, instruction='第一版最优'))
        plan.export()
        text = '\n'.join(p.text for p in Document(plan.folder / 'SOP.docx').paragraphs)
        self.assertIn('SOP待补', text)
        self.assertFalse((plan.folder / 'sop/v0001.docx').exists())


def demo(root):
    root = Path(root).resolve()
    if (root / 'report').exists() or (root / 'plans').exists():
        raise ValueError('Demo root must not already contain reports/plans')
    plan = PlanStore(root, 'DEMO', 'peg')
    plan.apply('create', request('init-demo', stage_name='示例工序', content=complete(),
                               source={'description': '用于软件版式验证的合成示例'}))
    plan.apply('set-best', request('choose-demo', version=1, instruction='验证流程：指定第一版为最优'))
    rows = [material(properties={'molar_mass': prop(100, 'g/mol')}),
            material(kind='solution', quantity='volume', value='10.0', unit='mL', component='测试溶质',
                     name='测试母液', properties={'molarity': prop('.1', 'mol/L'), 'molar_mass': prop(100, 'g/mol')}),
            material(kind='dispersion', quantity='volume', value='5.0', unit='mL', component='测试颗粒',
                     name='测试分散液', properties={'mass_concentration': prop(50, 'mg/mL')}),
            material(kind='unknown', value='0.80', name='测试产物', role='产物',
                     sample_id='DEMO-20260920-1', sample_step='S01')]
    item = daily(); item['materials'] = rows
    path = Store(root).save([item])
    output = {'daily': str(path), 'plans': plan.export(), 'fixture': 'synthetic software validation only'}
    (root / 'manifest.json').write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    import sys
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    if len(sys.argv) == 3 and sys.argv[1] == '--demo-root':
        demo(sys.argv[2])
    else:
        unittest.main()
