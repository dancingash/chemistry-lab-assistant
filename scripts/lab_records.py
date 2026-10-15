"""Local daily report storage. Run --help for scoped queries, saves and monthly trackers."""
from __future__ import annotations
from lab_config import configure_console, data_root

import argparse
from copy import deepcopy
import calendar
import hashlib
import json
import os
import re
import sys
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from extract_daily_docx import FIELD_NAMES, SAMPLE_RE, extract_docx

LEGACY_TRACKER = '样品进度追踪表.docx'


def tracker_name(month):
    y, m = month.split('-')
    return f'项目追踪表_{y}年{m}月.docx'
HEADERS = ['样品编号', '项目', '创建日期', '当前步骤与状态', '最近实验日期', '记录入口']
STATES = {'合成中', '待后处理', '后处理中', '待测试', '测试中', '已完成', '方案终止'}
COLORS = {'合成': ('2563EB', 'EFF6FF'), '后处理': ('16A34A', 'F0FDF4'), '测试': ('EA580C', 'FFF7ED')}
META = ('record_id', 'date', 'sample_id', 'sample_ids', 'project', 'created_date', 'step',
        'source_sample', 'source_step', 'experiment_type', 'state', 'origin_record_id',
        'plan_stage', 'plan_id', 'plan_version', 'batch_id', 'batch_revision', 'batch_group')


class Conflict(ValueError):
    pass


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def atomic_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=path.parent, suffix='.tmp')
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as stream:
            json.dump(data, stream, ensure_ascii=False, indent=2)
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def atomic_docx(doc, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=path.parent, suffix='.docx')
    os.close(fd)
    try:
        doc.save(name)
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def scoped(root, relative):
    # Historical Windows relative paths also work after data moves to macOS/Linux.
    path = (root / str(relative).replace('\\', '/')).resolve()
    if not path.is_relative_to(root.resolve()):
        raise Conflict('文件必须位于 数据根目录内')
    return path


def blank_document(title, landscape=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    sec.left_margin = sec.right_margin = Cm(1.6)
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Microsoft YaHei', Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.paragraph_format.space_after = Pt(5)
    doc.add_heading(title, 0)
    return doc


def record_text(record):
    y, m, d = record['date'].split('-')
    weekday = '一二三四五六日'[date.fromisoformat(record['date']).weekday()]
    lines = [f'{y}年{m}月{d}日 · 周{weekday}']
    if record.get('raw'):
        # Keep manually entered fields and migrated legacy text verbatim.
        return '\n'.join(lines) + '\n' + record['raw']
    for key, label in FIELD_NAMES.items():
        if record.get(key):
            lines.append(f'【{label}】{record[key]}')
    return '\n'.join(lines)


def write_daily(path, records):
    doc = blank_document('实验日报 · ' + records[0]['date'])
    from lab_materials_docx import add_material_table, attach_material_payload
    material_payload = {}
    for record in records:
        table = doc.add_table(rows=1, cols=1)
        table.autofit = False
        table.columns[0].width = Cm(17.8)
        cell = table.cell(0, 0)
        cell.text = record_text(record)
        if record.get('materials'):
            material_payload[record['record_id']] = {
                'table_hash': add_material_table(cell, record['materials']), 'materials': record['materials']}
        experiment_type = next((kind for kind in COLORS if record['experiment_type'].startswith(kind)), None)
        border, fill = COLORS[experiment_type or '合成']
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), fill)
        properties.append(shading)
        borders = OxmlElement('w:tcBorders')
        for side in ('top', 'left', 'bottom', 'right'):
            element = OxmlElement('w:' + side)
            for key, value in [('val', 'single'), ('sz', '8'), ('color', border)]:
                element.set(qn('w:' + key), value)
            borders.append(element)
        properties.append(borders)
        # No cantSplit: long experiment records can continue on the next page.
        doc.add_paragraph()
    attach_material_payload(doc, material_payload)
    atomic_docx(doc, path)


def tracker_data(path):
    if not path.exists():
        return None
    doc = Document(path)
    month = next((m[1] for p in doc.paragraphs if (m := re.search(r'维护月份：(\d{4}-\d{2})', p.text))), None)
    if not month or not doc.tables:
        raise Conflict(f'追踪表月份或表格结构不明确：{path}')
    table = doc.tables[0]
    if [c.text for c in table.rows[0].cells] != HEADERS:
        raise Conflict('追踪表列结构已改变，请核对后再同步')
    rows = [[c.text for c in row.cells] for row in table.rows[1:]]
    if len({r[0] for r in rows}) != len(rows):
        raise Conflict('追踪表中有重复样品行')
    return {'month': month, 'rows': rows}


def fingerprint(record):
    # Ignore storage location: copying the same record must not double count it.
    return hashlib.sha256(json.dumps({k: v for k, v in record.items() if k not in ('path',)},
                                     sort_keys=True, ensure_ascii=False).encode()).hexdigest()


class Store:
    def __init__(self, root):
        self.root = Path(root).resolve()
        self.report = self.root / 'report'
        self.index_path = self.report / 'index.json'
        self.log = []
        if self.index_path.exists():
            try:
                self.index = json.loads(self.index_path.read_text(encoding='utf-8'))
            except (ValueError, OSError) as exc:
                raise Conflict('索引损坏；使用 rebuild 从日报重建，旧日报需重新指定权威源') from exc
        else:
            self.index = {'version': 1, 'files': {}, 'tracker': None}
        if self.index.get('version') != 1:
            raise Conflict('不支持的索引版本')

    def parse(self, path):
        self.log.append({'action': 'read_docx', 'path': path.relative_to(self.root).as_posix()})
        return extract_docx(path)['records']

    def all_meta(self):
        origins = {r.get('origin_record_id') for item in self.index['files'].values()
                   if not item.get('legacy') for r in item['records'] if r.get('origin_record_id')}
        return [dict(r, path=p) for p, item in self.index['files'].items() for r in item['records']
                if not (item.get('legacy') and r['record_id'] in origins)]

    def current_tracker_path(self):
        old = self.index.get('tracker')
        if old and old.get('data'):
            month = old['data']['month']
            monthly = self.report / month / tracker_name(month)
            if monthly.exists():
                return monthly
        candidates = sorted(self.report.glob('????-??/项目追踪表_????年??月.docx'))
        if candidates:
            return candidates[-1]
        legacy = self.report / LEGACY_TRACKER
        return legacy if legacy.exists() else None

    def refresh(self, legacy=()):
        paths = {p.resolve() for p in self.report.glob('????-??/日报_????-??-??.docx')}
        paths.update(scoped(self.root, p) for p in self.index['files'])
        paths.update(scoped(self.root, p) for p in legacy)
        explicit = {str(scoped(self.root, p).relative_to(self.root)).replace('\\', '/') for p in legacy}
        # Only one chosen legacy snapshot per month. Never choose by filename or mtime.
        for p in sorted(paths):
            rel = p.relative_to(self.root).as_posix()
            if not p.exists():
                self.index['files'].pop(rel, None)
                continue
            stamp = digest(p)
            previous = self.index['files'].get(rel)
            if previous and previous['sha256'] == stamp:
                continue
            records = self.parse(p)
            daily = bool(re.fullmatch(r'日报_\d{4}-\d{2}-\d{2}\.docx', p.name))
            if not daily and rel in explicit:
                months = {r['date'][:7] for r in records if r['date']}
                for other, item in self.index['files'].items():
                    if other != rel and item.get('legacy') and months.intersection(item.get('months', [])):
                        raise Conflict(f'已有同月历史权威源 {other}；请用 replace-legacy 显式替换，不能合并多个快照')
            self.index['files'][rel] = {
                'sha256': stamp, 'legacy': not daily,
                'months': sorted({r['date'][:7] for r in records if r['date']}),
                'records': [dict({k: r.get(k, '') for k in META}, fingerprint=fingerprint(r)) for r in records],
            }
        seen = {}
        for r in self.all_meta():
            rid = r['record_id']
            if rid in seen and seen[rid] != r['fingerprint']:
                raise Conflict(f'记录标识冲突：{rid}；内容不同，不可静默去重')
            seen[rid] = r['fingerprint']
        origins = {}
        for r in self.all_meta():
            origin = r.get('origin_record_id')
            if origin:
                if origin in origins and origins[origin] != r['record_id']:
                    raise Conflict(f'同一旧记录被迁入多个每日卡片：{origin}')
                origins[origin] = r['record_id']
        self.check_tracker()
        atomic_json(self.index_path, self.index)
        # Daily Word is authoritative. Retrying refresh also recovers interrupted reverse links.
        from lab_plans import sync_record_links
        sync_record_links(self.root, self.all_meta())

    def check_tracker(self):
        path = self.current_tracker_path()
        if path is None:
            return
        current = tracker_data(path)
        old = self.index.get('tracker')
        if old and old['sha256'] == digest(path):
            return
        if old and old['data'] != current:
            raise Conflict('追踪表被手动修改；请对照日报核对。确认后的内容写入对应日报，再用 tracker --accept-reviewed 同步')
        # A missing index must not authorize discarding handwritten tracker rows.
        if not old and current['rows']:
            expected = self.rows_for_month(current['month'])
            if current['rows'] != [r['cells'] for r in expected]:
                raise Conflict('追踪表与可恢复的日报不一致；补接历史源或核对手动编辑后，再确认同步')
        self.index['tracker'] = {'sha256': digest(path), 'data': current}

    def query(self, sample=None, step=None, start=None, end=None, kind=None):
        matches = [r for r in self.all_meta()
                   if (not sample or sample in r['sample_ids'])
                   and (not step or r['step'] == step)
                   and (not kind or r['experiment_type'] == kind)
                   and (not start or (r['date'] and r['date'] >= start))
                   and (not end or (r['date'] and r['date'] <= end))]
        wanted = {r['record_id'] for r in matches}
        result = {}
        for p in sorted({r['path'] for r in matches}):
            for record in self.parse(scoped(self.root, p)):
                if record['record_id'] in wanted:
                    result.setdefault(record['record_id'], dict(record, path=p))
        return sorted(result.values(), key=lambda r: (r['date'] or '', r['record_id']))

    def rows_for_month(self, month):
        last = f'{month}-{calendar.monthrange(int(month[:4]), int(month[5:]))[1]:02d}'
        groups = {}
        unique = {}
        for record in self.all_meta():
            unique.setdefault(record['record_id'], record)
        for r in unique.values():
            if r['date'] and r['date'] <= last:
                for sample in r['sample_ids']:
                    groups.setdefault(sample, []).append(r)
        rows = []
        for sample, records in groups.items():
            records.sort(key=lambda r: r['date'])
            # Explicit states are needed; do not infer completion from a test result.
            states = [r for r in records if r['state']]
            if not states:
                # Old reports often omitted explicit states. Preserve uncertainty rather than
                # blocking unrelated new experiments or pretending old samples are unfinished.
                if not any(r['date'][:7] == month for r in records):
                    continue
                states = [dict(records[-1], state='待确认')]
            # A result measured on S01 cannot move an in-progress S02 back to S01.
            highest_step = max((int(r['step'][1:]) for r in states if re.fullmatch(r'S\d+', r['step'])), default=0)
            current_states = [r for r in states if r['step'] == f'S{highest_step:02d}'] if highest_step else states
            newest = current_states[-1]
            if newest['date'][:7] < month and newest['state'] in ('已完成', '方案终止'):
                continue
            created = next((r['created_date'] for r in records if r['created_date']), '')
            if not created:
                match = re.search(r'-(\d{4})(\d{2})(\d{2})-\d+', sample)
                created = '-'.join(match.groups()) if match else records[0]['date']
            links = [(r['date'] + ' ' + (r['step'] or '待补'), r['path'], r['record_id']) for r in records]
            cells = [sample, newest['project'], created,
                     (newest['step'] or '待补') + ' / ' + newest['state'], records[-1]['date'],
                     '\n'.join(label for label, _, _ in links)]
            rows.append({'cells': cells, 'links': links})
        def key(row):
            sample = row['cells'][0]
            match = re.search(r'-\d{8}-(\d+)', sample)
            return row['cells'][2], int(match[1]) if match else 0, sample
        return sorted(rows, key=key, reverse=True)

    def write_tracker(self, month, accept_reviewed=False):
        date.fromisoformat(month + '-01')
        current_path = self.current_tracker_path()
        current = tracker_data(current_path) if current_path else None
        target = self.report / month / tracker_name(month)
        if not accept_reviewed:
            self.check_tracker()
        rows = self.rows_for_month(month)
        if current and current['month'] > month:
            raise Conflict('不可将当前追踪表倒退到旧月份；请查询历史月份记录')
        if current and current['month'] < month and current_path == self.report / LEGACY_TRACKER:
            archive = self.report / current['month'] / tracker_name(current['month'])
            archive.parent.mkdir(parents=True, exist_ok=True)
            if archive.exists() and digest(archive) != digest(current_path):
                raise Conflict('上月追踪表已存在且内容不同，请核对')
            if not archive.exists():
                archive.write_bytes(current_path.read_bytes())
        doc = blank_document('项目追踪表', landscape=True)
        doc.add_paragraph('维护月份：' + month)
        table = doc.add_table(rows=1, cols=len(HEADERS))
        table.style = 'Table Grid'
        for c, text in zip(table.rows[0].cells, HEADERS):
            c.text = text
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, row['cells']):
                cell.text = text
            cells[-1].text = ''
            for i, (label, path, _) in enumerate(row['links']):
                p = cells[-1].paragraphs[0] if i == 0 else cells[-1].add_paragraph()
                link = OxmlElement('w:hyperlink')
                relative = os.path.relpath(scoped(self.root, path), target.parent).replace('\\', '/')
                link.set(qn('r:id'), doc.part.relate_to(relative, RT.HYPERLINK, is_external=True))
                run = OxmlElement('w:r')
                props = OxmlElement('w:rPr')
                color = OxmlElement('w:color'); color.set(qn('w:val'), '2563EB'); props.append(color)
                run.append(props)
                text = OxmlElement('w:t'); text.text = label; run.append(text)
                link.append(run); p._p.append(link)
        atomic_docx(doc, target)
        self.index['tracker'] = {'sha256': digest(target), 'data': tracker_data(target)}
        atomic_json(self.index_path, self.index)
        return target

    def migrate_day(self, day, legacy_path):
        date.fromisoformat(day)
        path = scoped(self.root, legacy_path)
        rel = path.relative_to(self.root).as_posix()
        if not self.index['files'].get(rel, {}).get('legacy'):
            raise Conflict('请先选择并登记该月权威旧日报，再迁入每日卡片')
        source = [r for r in self.parse(path) if r['date'] == day]
        if not source:
            raise Conflict('旧日报中没有该日记录')
        daily = self.report / day[:7] / f'日报_{day}.docx'
        if daily.exists():
            existing = self.parse(daily)
            if {r['record_id'] for r in source}.issubset({r['origin_record_id'] for r in existing}):
                return daily
            raise Conflict('该日卡片已经存在且内容不同；请先核对再迁移')
        cards = []
        for r in source:
            copied = dict(r, record_id='card-' + r['record_id'], origin_record_id=r['record_id'])
            copied['raw'] = f"【记录标识】{copied['record_id']}\n【来源记录标识】{r['record_id']}\n" + r['raw']
            cards.append(copied)
        write_daily(daily, cards)
        self.refresh()
        current = tracker_data(self.current_tracker_path()) if self.current_tracker_path() else None
        self.write_tracker(max(day[:7], current['month'] if current else day[:7]))
        return daily

    def save(self, records):
        records = deepcopy(records)
        if not records:
            raise Conflict('没有待写入记录')
        self.refresh()
        existing_meta = self.all_meta()
        allowed_fields = set(FIELD_NAMES) | {'date', 'materials', 'reviewed_materials'}
        for r in records:
            if set(r) - allowed_fields:
                raise Conflict('输入包含未知字段：' + ', '.join(sorted(set(r) - allowed_fields)))
            if any(not isinstance(value, str) for key, value in r.items() if key != 'materials'):
                raise Conflict('记录字段必须是字符串')
            if 'materials' in r:
                from lab_quantities import prepare_materials
                r['materials'] = prepare_materials(r['materials'])
            for key in ('record_id', 'date', 'project', 'sample_id', 'created_date', 'step', 'experiment_type', 'state', 'today', 'follow_up'):
                if not isinstance(r.get(key), str) or not r[key].strip():
                    raise Conflict(f'缺少字段：{key}')
            date.fromisoformat(r['date'])
            from lab_plans import validate_record_plan
            validate_record_plan(self.root, r)
            if any(r.get(k) for k in ('batch_id', 'batch_revision', 'batch_group')):
                from lab_formulations import validate_record_batch
                validate_record_batch(self.root, r)
            pending_identity = not SAMPLE_RE.fullmatch(r['sample_id'])
            if pending_identity and (r['created_date'], r['step'], r['state']) != ('待确认', '待确认', '待确认'):
                raise Conflict('样品身份未确定时，创建日期、步骤和状态均须标为待确认')
            if not pending_identity:
                date.fromisoformat(r['created_date'])
            if not pending_identity and r['created_date'] > r['date']:
                raise Conflict('创建日期不能晚于实验日期')
            if not pending_identity and not re.fullmatch(r'S\d{2,}', r['step']):
                raise Conflict('样品编号或步骤格式不合法')
            if (not pending_identity and r['state'] not in STATES) or r['experiment_type'] not in COLORS:
                raise Conflict('未知状态或实验类型')
            if '-B' in r['sample_id'] and not (r.get('source_sample') and r.get('source_step')):
                raise Conflict('分支样品必须注明来源样品和来源步骤')
            if bool(r.get('source_sample')) != bool(r.get('source_step')):
                raise Conflict('来源样品和来源步骤必须同时提供')
            if r.get('source_sample'):
                if not SAMPLE_RE.fullmatch(r['source_sample']) or not re.fullmatch(r'S\d{2,}', r['source_step']):
                    raise Conflict('来源样品或步骤格式不合法')
                if not any(old['sample_id'] == r['source_sample'] and old['step'] == r['source_step'] for old in existing_meta + records):
                    raise Conflict('来源步骤未找到；请先接入来源记录')
                if r['source_sample'] == r['sample_id'] and r['source_step'] == r['step']:
                    raise Conflict('记录不能引用自己的当前步骤作为来源')
            for old in existing_meta:
                if old['record_id'] == r['record_id'] and old['date'] != r['date']:
                    raise Conflict('修改记录日期需显式迁移，不能复制同一标识到另一天')
                if old['record_id'] == r['record_id'] and self.index['files'][old['path']]['legacy']:
                    raise Conflict('该标识属于历史权威源；请定点修改原记录，不能另存副本造成双重来源')
                if old['sample_id'] == r['sample_id'] and old['created_date'] and old['created_date'] != r['created_date']:
                    raise Conflict('同一样品的创建日期不能改变')
            if any(other['sample_id'] == r['sample_id'] and other['created_date'] != r['created_date'] for other in records):
                raise Conflict('同批输入的样品创建日期不一致')
        if len({r['record_id'] for r in records}) != len(records):
            raise Conflict('输入中存在重复记录标识')
        if len({r['date'] for r in records}) != 1:
            raise Conflict('一次 save 只写一个日期；跨日记录分别提交')
        day = records[0]['date']
        path = self.report / day[:7] / f'日报_{day}.docx'
        existing = self.parse(path) if path.exists() else []
        if path.exists():
            old_doc = Document(path)
            extra = [p.text for p in old_doc.paragraphs if p.text.strip() and p.text != '实验日报 · ' + day]
            if extra or len(old_doc.tables) != len(existing) or old_doc.inline_shapes or old_doc.element.xpath('.//w:drawing'):
                raise Conflict('当日日报含手动正文或图片，请使用定点编辑保留内容')
        # Rebuild only known managed cards; prevent losing legacy/manual unknown fields.
        for r in existing:
            if r['record_id'].startswith('legacy-'):
                raise Conflict('当日日报存在无标识记录，请先核对格式，避免覆盖手动内容')
        by_id = {r['record_id']: r for r in existing}
        for r in records:
            old = by_id.get(r['record_id'])
            if old and (set(old['fields']) - set(FIELD_NAMES.values())):
                raise Conflict('待修改卡片含自定义字段；请定点编辑该卡片以保留内容')
            if old and old.get('materials'):
                from lab_quantities import raw_materials, quantities_hash
                before = raw_materials(old['materials'])
                after = raw_materials(r.get('materials', old['materials']))
                if before != after and r.get('reviewed_materials') != quantities_hash(before):
                    raise Conflict('同一记录用量冲突，核对后以 reviewed_materials 确认旧数据摘要：' +
                                   quantities_hash(before) + '\n' + json.dumps({'before': before, 'after': after}, ensure_ascii=False))
                if 'materials' not in r:
                    r['materials'] = old['materials']
            r.pop('reviewed_materials', None)
            by_id[r['record_id']] = r
        write_daily(path, list(by_id.values()))
        # A failure here leaves the daily Word authoritative. `refresh` + `tracker` recover.
        self.refresh()
        current_path = self.current_tracker_path()
        current = tracker_data(current_path) if current_path else None
        self.write_tracker(max(day[:7], current['month'] if current else day[:7]))
        return path


def main():
    configure_console()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=None)
    sub = parser.add_subparsers(dest='command', required=True)
    refresh = sub.add_parser('refresh')
    refresh.add_argument('--legacy', action='append', default=[], help='Explicitly chosen authoritative legacy Word, relative to root')
    rebuild = sub.add_parser('rebuild')
    rebuild.add_argument('--legacy', action='append', default=[])
    query = sub.add_parser('query')
    for flag in ('sample', 'step', 'start', 'end', 'kind'):
        query.add_argument('--' + flag)
    query.add_argument('--out', type=Path)
    query.add_argument('--full', action='store_true', help='Include duplicated raw text/label fields for parser diagnostics')
    save = sub.add_parser('save'); save.add_argument('input', type=Path)
    tracker = sub.add_parser('tracker'); tracker.add_argument('--month', required=True)
    tracker.add_argument('--accept-reviewed', action='store_true', help='Only after resolving manual tracker changes against daily records')
    replace = sub.add_parser('replace-legacy')
    replace.add_argument('--old', required=True); replace.add_argument('--new', required=True)
    migrate = sub.add_parser('migrate-day')
    migrate.add_argument('--date', required=True); migrate.add_argument('--legacy', required=True)
    args = parser.parse_args()
    args.root = data_root(args.root)
    try:
        if args.command == 'rebuild':
            # Keep the old index intact until the new index has passed validation.
            store = Store.__new__(Store)
            store.root = args.root.resolve(); store.report = store.root / 'report'
            store.index_path = store.report / 'index.json'; store.log = []
            store.index = {'version': 1, 'files': {}, 'tracker': None}
        else:
            store = Store(args.root)
        output = {}
        if args.command in ('refresh', 'rebuild'):
            store.refresh(args.legacy)
        elif args.command == 'replace-legacy':
            old = scoped(store.root, args.old).relative_to(store.root).as_posix()
            if not store.index['files'].get(old, {}).get('legacy'):
                raise Conflict('old 必须是已登记的历史权威源')
            store.index['files'].pop(old)
            store.refresh([args.new])
        elif args.command == 'query':
            if not (args.sample or args.start or args.end):
                raise Conflict('查询必须指定样品或日期范围，避免全量返回')
            for value in (args.start, args.end):
                if value:
                    date.fromisoformat(value)
            if args.start and args.end and args.start > args.end:
                raise Conflict('查询起始日期晚于结束日期')
            store.refresh()
            output['records'] = store.query(args.sample, args.step, args.start, args.end, args.kind)
            if not args.full:
                for record in output['records']:
                    record.pop('raw', None)
                    # Keep unknown historical fields, but avoid repeating all known content.
                    record['fields'] = {k: v for k, v in record['fields'].items() if k not in FIELD_NAMES.values()}
        elif args.command == 'save':
            data = json.loads(args.input.read_text(encoding='utf-8-sig'))
            output['path'] = str(store.save(data['records']))
            from project_context import route_status
            projects = sorted({r['project'].split()[0] for r in data['records']})
            output['route_checks'] = [route_status(args.root, project) for project in projects]
        elif args.command == 'tracker':
            if args.accept_reviewed:
                # Explicit reconciliation still requires the daily facts to be complete.
                store.index['tracker'] = None
                current_path = store.current_tracker_path()
                current = tracker_data(current_path) if current_path else None
                if current:
                    store.index['tracker'] = {'sha256': digest(current_path), 'data': current}
            store.refresh()
            output['path'] = str(store.write_tracker(args.month, args.accept_reviewed))
        elif args.command == 'migrate-day':
            output['path'] = str(store.migrate_day(args.date, args.legacy))
        output['read_log'] = store.log
        if getattr(args, 'out', None):
            atomic_json(args.out, output)
        print(json.dumps(output, ensure_ascii=False, indent=2))
    except (ValueError, OSError) as exc:
        print(json.dumps({'error': str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
