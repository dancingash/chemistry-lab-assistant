import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from install import install, package_files
from lab_config import data_root
from doctor import check
from demo import create_demo


class PortableInstallTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_clean_copy_installs_without_original_and_has_no_private_content(self):
        source = self.folder / 'clone 中文 space'
        source.mkdir()
        # Reproduce only the complete declared release, not the local workspace.
        for relative, path in package_files(ROOT).items():
            if not relative.startswith(('scripts/', 'templates/', 'examples/')) and relative not in ('README.md', 'VALIDATION.md', 'requirements.lock', 'pyproject.toml'):
                relative = 'skills/lab-rd/' + relative
            target = source / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, target)
        destination = self.folder / 'skills 中文'
        data = self.folder / 'user data'
        install(source, destination, data, skip_deps=True)
        shutil.move(str(source), str(self.folder / 'removed-clone'))
        package = destination / 'lab-rd'
        env = dict(os.environ)
        env.pop('LAB_RD_DATA_ROOT', None)
        env.pop('REPORT_SKILLS_ROOT', None)
        result = subprocess.run([sys.executable, str(package / 'scripts/lab_records.py'),
            'query', '--start', '2030-01-01', '--end', '2030-01-31'], env=env,
            check=True, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(json.loads(result.stdout)['records'], [])
        self.assertTrue((data / 'report/index.json').exists())
        self.assertFalse((package / 'report').exists())
        for relative in json.loads((package / 'installation.json').read_text())['files']:
            path = package / relative
            if path.suffix in ('.py', '.md', '.json', '.yaml', '.mjs', '.toml', '.lock'):
                text = path.read_text(encoding='utf-8')
                for forbidden in (str(ROOT), str(Path.home())):
                    self.assertNotIn(forbidden, text, relative)
        with ZipFile(package / 'templates/monthly-summary.xlsx') as archive:
            text = ''.join(archive.read(n).decode('utf-8', errors='ignore') for n in archive.namelist() if n.endswith('.xml'))
            for forbidden in (str(ROOT), str(Path.home())):
                self.assertNotIn(forbidden, text)
            from openpyxl import load_workbook
            workbook = load_workbook(package / 'templates/monthly-summary.xlsx', read_only=True)
            self.assertEqual(workbook.sheetnames, ['Template'])
            sheet = workbook.active
            self.assertTrue(all(sheet.cell(row, col).value is None for row in range(3, 7) for col in range(2, 6)))
            workbook.close()

    def test_repeat_upgrade_preserves_data_config_custom_files_and_backups(self):
        skills = self.folder / 'skills'
        data = self.folder / 'data'
        install(ROOT, skills, data, skip_deps=True)
        package = skills / 'lab-rd'
        (data / 'report/fact.txt').write_text('original fact')
        (package / 'custom.md').write_text('user custom')
        config = json.loads((package / 'local.json').read_text())
        config['user_key'] = 'keep'
        (package / 'local.json').write_text(json.dumps(config))
        (package / 'SKILL.md').write_text('user edited skill')
        result = install(ROOT, skills, skip_deps=True)
        self.assertEqual((data / 'report/fact.txt').read_text(), 'original fact')
        self.assertEqual((package / 'custom.md').read_text(), 'user custom')
        self.assertEqual(json.loads((package / 'local.json').read_text())['user_key'], 'keep')
        self.assertEqual(Path(result['backup']).joinpath('lab-rd/SKILL.md').read_text(), 'user edited skill')
        self.assertEqual(result['data_dir'], str(data.resolve()))

    def test_two_users_have_separate_data(self):
        for user in ('a', 'b'):
            install(ROOT, self.folder / user / 'skills', self.folder / user / 'data', skip_deps=True)
        (self.folder / 'a/data/references/project.yaml').write_text('private')
        self.assertFalse((self.folder / 'b/data/references/project.yaml').exists())

    def test_path_precedence(self):
        with patch.dict(os.environ, {'LAB_RD_DATA_ROOT': str(self.folder/'new'), 'REPORT_SKILLS_ROOT': str(self.folder/'old')}), patch('lab_config.read_config', return_value={'data_dir': str(self.folder/'config')}):
            self.assertEqual(data_root(self.folder/'explicit'), (self.folder/'explicit').resolve())
            self.assertEqual(data_root(), (self.folder/'new').resolve())
            with patch.dict(os.environ, {'LAB_RD_DATA_ROOT': ''}):
                self.assertEqual(data_root(), (self.folder/'old').resolve())
                with patch.dict(os.environ, {'REPORT_SKILLS_ROOT': ''}):
                    self.assertEqual(data_root(), (self.folder/'config').resolve())

    def test_missing_capabilities_do_not_claim_full_readiness(self):
        result = check()
        self.assertFalse(result['full_report_ready'])
        with patch('doctor.importlib.import_module', side_effect=ImportError('missing')):
            result = check()
            self.assertFalse(result['core_ready'])
            self.assertTrue(any('reinstall' in issue or 'rerun' in issue for issue in result['issues']))

    def test_install_rejects_data_inside_program_and_unmanaged_targets(self):
        skills = self.folder / 'skills'
        with self.assertRaises(ValueError):
            install(ROOT, skills, skills/'lab-rd/data', skip_deps=True)
        (skills/'lab-rd').mkdir(parents=True)
        (skills/'lab-rd/custom').write_text('keep')
        with self.assertRaises(ValueError):
            install(ROOT, skills, self.folder/'data', skip_deps=True)

    def test_demo_links_two_days_to_one_plan_and_refuses_overwrite(self):
        root = self.folder / 'fictional-demo'
        with patch('demo.subprocess.run', wraps=subprocess.run):
            create_demo(root)
        from lab_records import Store
        records = Store(root).query(sample='DEMO-20300107-1')
        self.assertEqual(len(records), 2)
        self.assertEqual({r['plan_version'] for r in records}, {'1'})
        plan = json.loads((root/'plans/DEMO/validation/plan.json').read_text(encoding='utf-8'))
        self.assertIsNone(plan['best_version'])
        with self.assertRaises(ValueError):
            create_demo(root)

    def test_dependency_failure_can_retry_and_does_not_retire_legacy(self):
        skills = self.folder/'skills'
        legacy = skills/'lab-daily-report'
        legacy.mkdir(parents=True)
        (legacy/'SKILL.md').write_text('user customized original')
        with patch('install.importlib.util.find_spec', return_value=object()), patch('install.venv.EnvBuilder.create', side_effect=OSError('test dependency failure')):
            with self.assertRaises(OSError):
                install(ROOT, skills, self.folder/'data', retire_legacy=True)
        self.assertEqual((legacy/'SKILL.md').read_text(), 'user customized original')
        result = install(ROOT, skills, skip_deps=True)
        self.assertEqual(result['data_dir'], str((self.folder/'data').resolve()))
        self.assertFalse((skills/'lab-rd/.installation-pending.json').exists())

    def test_windows_relative_path_is_portable_and_traversal_is_blocked(self):
        from lab_records import scoped, Conflict
        file = self.folder/'report/中文 daily.docx'
        file.parent.mkdir()
        file.write_text('test')
        self.assertEqual(scoped(self.folder, 'report\\中文 daily.docx'), file.resolve())
        with self.assertRaises(Conflict):
            scoped(self.folder, '..\\outside.docx')


if __name__ == '__main__':
    unittest.main()
